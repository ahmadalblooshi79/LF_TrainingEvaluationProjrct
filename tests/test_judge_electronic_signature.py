import io
import unittest
import zipfile
from datetime import datetime
from pathlib import Path
from tempfile import TemporaryDirectory

from PIL import Image
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker

from app.database import Base
from app.evaluation_workflow import apply_judge_approve, apply_other_judge_overwrite
from app.judge_signature import (
    JudgeSignatureError,
    attach_signature_snapshot,
    delete_master_for_user_ids,
    delete_signatures_for_exercise_judges,
    get_master,
    make_sample_signature_png,
    png_is_rgba,
    resolve_approval_png,
    save_or_replace_master,
    snapshot_png,
    validate_transparent_signature_png,
)
from app.models import (
    EvaluationListPdfItem,
    EvaluationListSavedResult,
    PlannerFlowBundleEvalSavedResult,
    User,
)


class _DummyUser:
    def __init__(self, user_id: int):
        self.id = user_id


class JudgeElectronicSignatureTests(unittest.TestCase):
    def setUp(self):
        engine = create_engine("sqlite:///:memory:")
        Base.metadata.create_all(engine)
        self.db = sessionmaker(bind=engine)()
        self.judge_a = User(
            username="judge_a",
            full_name="أحمد",
            password_hash="x",
            role_key="judge",
        )
        self.judge_b = User(
            username="judge_b",
            full_name="سالم",
            password_hash="x",
            role_key="judge",
        )
        self.db.add_all([self.judge_a, self.judge_b])
        self.db.flush()
        self.png = make_sample_signature_png()

    def tearDown(self):
        self.db.close()

    def test_sample_png_has_alpha_not_white_background(self):
        self.assertTrue(png_is_rgba(self.png))
        validate_transparent_signature_png(self.png)
        im = Image.open(io.BytesIO(self.png))
        self.assertEqual(im.mode, "RGBA")
        corners = [im.getpixel((0, 0)), im.getpixel((im.width - 1, im.height - 1))]
        for px in corners:
            self.assertLess(px[3], 16)

    def test_white_jpeg_rejected(self):
        im = Image.new("RGB", (80, 40), (255, 255, 255))
        buf = io.BytesIO()
        im.save(buf, format="JPEG")
        with self.assertRaises(JudgeSignatureError):
            validate_transparent_signature_png(buf.getvalue())

    def test_user_isolation(self):
        save_or_replace_master(self.db, int(self.judge_a.id), self.png, replace=True)
        self.assertIsNotNone(get_master(self.db, int(self.judge_a.id)))
        self.assertIsNone(get_master(self.db, int(self.judge_b.id)))
        with self.assertRaises(JudgeSignatureError) as ctx:
            resolve_approval_png(
                self.db,
                _DummyUser(int(self.judge_b.id)),
                request_user_id=int(self.judge_a.id),
            )
        self.assertEqual(ctx.exception.code, "owner_mismatch")

    def test_no_signature_blocks_approval_resolution(self):
        with self.assertRaises(JudgeSignatureError) as ctx:
            resolve_approval_png(self.db, _DummyUser(int(self.judge_a.id)))
        self.assertEqual(ctx.exception.code, "no_signature")

    def test_historical_snapshot_survives_master_update(self):
        v1 = make_sample_signature_png(width=200, height=70)
        save_or_replace_master(self.db, int(self.judge_a.id), v1, replace=True)
        item = EvaluationListPdfItem(
            exercise_id=1,
            exercise_phase="preparation",
            unit_level_key="ul_x",
            unit_level_label="x",
            sort_order=0,
            text="قائمة أ",
            pdf_relpath="x.xlsx",
        )
        self.db.add(item)
        self.db.flush()
        saved = EvaluationListSavedResult(
            evaluation_item_id=int(item.id),
            exercise_id=1,
            payload_json='{"rows":[]}',
        )
        self.db.add(saved)
        self.db.flush()
        apply_judge_approve(saved, int(self.judge_a.id))
        attach_signature_snapshot(
            saved, user_id=int(self.judge_a.id), png_bytes=v1, version=1
        )
        v2 = make_sample_signature_png(width=240, height=90)
        save_or_replace_master(self.db, int(self.judge_a.id), v2, replace=True)
        self.assertEqual(snapshot_png(saved), v1)
        master = get_master(self.db, int(self.judge_a.id))
        self.assertEqual(int(master.version), 2)
        self.assertEqual(bytes(master.png_blob), v2)

    def test_planner_saved_result_also_stores_snapshot(self):
        from app.models.domain import (
            ExercisePlannerFlowBundle,
            ExercisePlannerFlowBundleActionEval,
        )

        bundle = ExercisePlannerFlowBundle(
            exercise_id=1,
            exercise_phase="preparation",
            unit_level_key="ul_x",
            unit_level_label="x",
        )
        self.db.add(bundle)
        self.db.flush()
        slot = ExercisePlannerFlowBundleActionEval(
            bundle_id=int(bundle.id),
            slot_index=1,
            title="إجراء",
            file_relpath="a.xlsx",
        )
        self.db.add(slot)
        self.db.flush()
        saved = PlannerFlowBundleEvalSavedResult(
            bundle_action_eval_id=int(slot.id),
            exercise_id=1,
            payload_json='{"rows":[]}',
        )
        self.db.add(saved)
        self.db.flush()
        apply_judge_approve(saved, int(self.judge_a.id))
        attach_signature_snapshot(
            saved, user_id=int(self.judge_a.id), png_bytes=self.png, version=1
        )
        self.assertEqual(snapshot_png(saved), self.png)

    def test_excel_export_keeps_png_alpha(self):
        from openpyxl import Workbook

        from app.evaluation_list_export import build_evaluation_list_xlsx_bytes

        with TemporaryDirectory() as td:
            src = Path(td) / "src.xlsx"
            wb = Workbook()
            ws = wb.active
            ws.title = "قائمة التقييم"
            ws["B1"] = "عنوان"
            ws["E20"] = "المحكم:"
            ws["G20"] = "أدخل اسم المحكم"
            ws["E21"] = "التوقيع:"
            wb.save(src)
            wb.close()
            data = build_evaluation_list_xlsx_bytes(
                src,
                doc_title="عنوان",
                unit_label="وحدة",
                date_str="2026-09-13",
                commander_name="قائد",
                judge_name="أحمد",
                eval_rows=[],
                saved_rows=[],
                signature_png=self.png,
                approved_at=datetime(2026, 9, 13, 10, 15),
            )
        from openpyxl import load_workbook

        wb2 = load_workbook(io.BytesIO(data))
        try:
            ws2 = wb2.active
            for row in ws2.iter_rows(values_only=True):
                for val in row:
                    if val is None:
                        continue
                    self.assertNotIn("اعتماد إلكتروني", str(val))
            imgs = list(getattr(ws2, "_images", None) or [])
            self.assertTrue(imgs)
            anchor = imgs[0].anchor
            from_cell = getattr(anchor, "_from", None)
            if from_cell is not None:
                self.assertEqual(int(from_cell.col), 6)
                self.assertEqual(int(from_cell.row) + 1, 21)
        finally:
            wb2.close()
        with zipfile.ZipFile(io.BytesIO(data)) as zf:
            media = [n for n in zf.namelist() if n.startswith("xl/media/")]
            self.assertTrue(media)
            png = zf.read(media[0])
        self.assertTrue(png_is_rgba(png))
        validate_transparent_signature_png(png)

    def test_excel_export_writes_acquired_marks(self):
        from openpyxl import Workbook, load_workbook

        from app.evaluation_list_export import build_evaluation_list_xlsx_bytes
        from app.evaluation_sheet_parser import read_evaluation_list_sheet

        with TemporaryDirectory() as td:
            src = Path(td) / "src.xlsx"
            wb = Workbook()
            ws = wb.active
            ws.title = "قائمة التقييم"
            ws["B1"] = "عنوان القائمة"
            ws["B2"] = "عناصــــــر التقييـــــم"
            ws["E2"] = "العلامـــــــات"
            ws["I2"] = "ملاحظــــــات"
            ws["E3"] = "القصوى"
            ws["F3"] = "المكتسبة"
            ws["G3"] = "النسبة"
            ws["H3"] = "النتيجة"
            ws["B4"] = "( أ )"
            ws["E4"] = "(ب)"
            ws["F4"] = "(جـ)"
            for i in range(1, 9):
                r = 4 + i
                ws[f"B{r}"] = f"{i}. بند"
                ws[f"E{r}"] = 5
                ws[f"F{r}"] = "أدخل العلامة"
                ws[f"G{r}"] = f"=IFERROR(F{r}/E{r},\"\")"
            ws["B13"] = "إجمالي العلامات"
            ws["E13"] = "=SUM(E5:E12)"
            ws["F13"] = "=SUM(F5:F12)"
            ws["E20"] = "المحكم:"
            ws["G20"] = "أدخل اسم المحكم"
            ws["E21"] = "التوقيع:"
            wb.save(src)
            wb.close()
            eval_rows = read_evaluation_list_sheet(src).get("eval_rows") or []
            saved = [
                {"acquired": "4", "notes": "", "row_kind": "score"}
                for _ in eval_rows
            ]
            data = build_evaluation_list_xlsx_bytes(
                src,
                doc_title="عنوان القائمة",
                unit_label="وحدة",
                date_str="2026-09-07",
                commander_name="قائد",
                judge_name="أحمد",
                eval_rows=eval_rows,
                saved_rows=saved,
            )
        wb2 = load_workbook(io.BytesIO(data))
        try:
            ws2 = wb2.active
            self.assertEqual(ws2["F3"].value, "المكتسبة")
            self.assertGreaterEqual(len(eval_rows), 8)
            self.assertEqual(ws2["F5"].value, 4)
            self.assertEqual(ws2["F7"].value, 4)
            self.assertEqual(ws2["G20"].value, "أحمد")
        finally:
            wb2.close()

    def test_other_judge_overwrite_clears_approval(self):
        item = EvaluationListPdfItem(
            exercise_id=1,
            exercise_phase="opening",
            unit_level_key="ul_x",
            unit_level_label="x",
            sort_order=0,
            text="قائمة",
            pdf_relpath="x.xlsx",
        )
        self.db.add(item)
        self.db.flush()
        saved = EvaluationListSavedResult(
            evaluation_item_id=int(item.id),
            exercise_id=1,
            payload_json='{"rows":[]}',
            saved_by_id=int(self.judge_a.id),
        )
        self.db.add(saved)
        self.db.flush()
        apply_judge_approve(saved, int(self.judge_a.id))
        attach_signature_snapshot(
            saved,
            user_id=int(self.judge_a.id),
            png_bytes=self.png,
            version=1,
        )
        apply_other_judge_overwrite(saved)
        self.assertFalse(saved.is_approved)
        self.assertIsNone(saved.approved_by_id)
        self.assertIsNone(snapshot_png(saved))

    def test_delete_signatures_for_exercise_judges(self):
        from app.models import (
            Exercise,
            ExerciseRosterKind,
            ExerciseRosterRow,
            ExerciseStatus,
        )

        owner = User(
            username="admin_sig_wipe",
            full_name="مدير",
            password_hash="x",
            role_key="system_admin",
        )
        self.db.add(owner)
        self.db.flush()
        ex = Exercise(
            code="T-SIG-1",
            title="تمرين تواقيع",
            owner_id=int(owner.id),
            status=ExerciseStatus.ACTIVE.value,
            exercise_type="trial",
            trained_unit="unit",
        )
        self.db.add(ex)
        self.db.flush()
        save_or_replace_master(self.db, int(self.judge_a.id), self.png, replace=True)
        save_or_replace_master(self.db, int(self.judge_b.id), self.png, replace=True)
        self.db.add(
            ExerciseRosterRow(
                exercise_id=int(ex.id),
                roster_kind=ExerciseRosterKind.JUDGE.value,
                military_number=self.judge_a.username,
                full_name=self.judge_a.full_name,
                unit_level_key="ul_x",
            )
        )
        self.db.commit()
        delete_signatures_for_exercise_judges(self.db, int(ex.id))
        self.db.commit()
        self.assertIsNone(get_master(self.db, int(self.judge_a.id)))
        self.assertIsNotNone(get_master(self.db, int(self.judge_b.id)))
        delete_master_for_user_ids(self.db, {int(self.judge_b.id)})
        self.db.commit()
        self.assertIsNone(get_master(self.db, int(self.judge_b.id)))


if __name__ == "__main__":
    unittest.main()

"""مركز التقارير: فلترة عنصر-بعنصر، عزل الإجراءات/المعاضل، وExcel."""
from __future__ import annotations

import json
import unittest
from datetime import datetime
from io import BytesIO
from types import SimpleNamespace
from unittest.mock import MagicMock, patch

from openpyxl import load_workbook
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker

from app.database import Base
from app.judge_eval_results_export import (
    KIND_PROCEDURE,
    POSITION_APPROVED,
    build_judge_results_xlsx,
    list_result_fields,
)
from app.models import (
    EvaluationListPdfItem,
    EvaluationListSavedResult,
    Exercise,
    ExercisePlannerFlowBundle,
    ExercisePlannerFlowBundleActionEval,
    ExerciseRosterKind,
    ExerciseRosterRow,
    ExerciseStatus,
    PlannerFlowBundleEvalSavedResult,
    User,
)
from app.models.user import RoleKey
from app.reports_center import (
    KIND_DILEMMA,
    KIND_PROCEDURE as REPORT_KIND_PROCEDURE,
    ORGANIZE_JUDGE,
    ORGANIZE_LIST,
    ORGANIZE_UNIT,
    RATING_ORDER,
    ReportFilters,
    ReportScope,
    collect_report_lists,
    explode_scored_items,
    filter_catalog,
    normalize_rating_label,
    parse_report_filters,
    query_report_items,
    rating_percentages,
    report_filename,
    run_report,
)
from app.reports_center_excel import build_detailed_report_xlsx as build_xlsx
from app.reports_center_excel import unique_sheet_name


def _score_row(element: str, acquired: float, max_val: float = 5, notes: str = "") -> dict:
    return {
        "row_kind": "score",
        "element": element,
        "acquired": acquired,
        "max_val": max_val,
        "notes": notes,
    }


class ReportsCenterTests(unittest.TestCase):
    def setUp(self):
        engine = create_engine("sqlite:///:memory:")
        Base.metadata.create_all(engine)
        self.db = sessionmaker(bind=engine)()
        self.judge = User(
            username="j_rep",
            password_hash="x",
            full_name="أحمد المحكم",
            role_key=RoleKey.JUDGE.value,
        )
        self.db.add(self.judge)
        self.db.flush()
        self.ex = Exercise(
            code="RPT-1",
            title="تمرين التقارير",
            owner_id=int(self.judge.id),
            status=ExerciseStatus.ACTIVE.value,
        )
        self.db.add(self.ex)
        self.db.flush()
        self.uk = "ul_alpha"
        self.db.add(
            ExerciseRosterRow(
                exercise_id=int(self.ex.id),
                roster_kind=ExerciseRosterKind.JUDGE.value,
                sort_order=0,
                military_number="j_rep",
                rank_ar="عقيد",
                full_name="أحمد المحكم",
                unit_level_key=self.uk,
                exercise_phase="preparation",
            )
        )
        self.item = EvaluationListPdfItem(
            exercise_id=int(self.ex.id),
            exercise_phase="preparation",
            unit_level_key=self.uk,
            unit_level_label="الوحدة أ",
            text="قائمة الإنقاذ",
            pdf_relpath="x.xlsx",
        )
        self.db.add(self.item)
        self.db.flush()
        rows = [
            {"row_kind": "section", "element": "عنوان"},
            _score_row("بند ممتاز", 5, notes="ملاحظة ممتاز"),
            _score_row("بند جيد جدا", 4.2),
            _score_row("بند جيد", 3.7),
            _score_row("بند مقبول 1", 3.2, notes="ضعف 1"),
            _score_row("بند مقبول 2", 3.1, notes="ضعف 2"),
            _score_row("بند مقبول 3", 3.0),
            _score_row("بند مقبول 4", 3.05, notes="ضعف 4"),
            _score_row("بند راسب", 2.0, notes="راسب"),
            {"row_kind": "score", "element": "لا ينطبق", "acquired": "na", "max_val": 5, "notes": ""},
        ]
        self.db.add(
            EvaluationListSavedResult(
                evaluation_item_id=int(self.item.id),
                exercise_id=int(self.ex.id),
                exercise_phase="preparation",
                unit_level_key=self.uk,
                payload_json=json.dumps({"rows": rows}, ensure_ascii=False),
                total_pct=70,
                grade_label="جيد",
                saved_by_id=int(self.judge.id),
                is_approved=True,
            )
        )
        bundle = ExercisePlannerFlowBundle(
            exercise_id=int(self.ex.id),
            exercise_phase="opening",
            unit_level_key=self.uk,
            unit_level_label="الوحدة أ",
        )
        self.db.add(bundle)
        self.db.flush()
        slot = ExercisePlannerFlowBundleActionEval(
            bundle_id=int(bundle.id),
            slot_index=1,
            title="قائمة معضلة كمين",
            file_relpath="d.xlsx",
        )
        self.db.add(slot)
        self.db.flush()
        self.slot = slot
        self.db.add(
            PlannerFlowBundleEvalSavedResult(
                bundle_action_eval_id=int(slot.id),
                exercise_id=int(self.ex.id),
                exercise_phase="opening",
                unit_level_key=self.uk,
                payload_json=json.dumps(
                    {
                        "rows": [
                            _score_row("عنصر معضلة ممتاز", 5),
                            _score_row("عنصر معضلة راسب", 1, notes="خلل"),
                        ]
                    },
                    ensure_ascii=False,
                ),
                total_pct=60,
                grade_label="مقبول",
                saved_by_id=int(self.judge.id),
                is_approved=False,
            )
        )
        self.db.commit()
        self.proc_scope = ReportScope(
            exercise_id=int(self.ex.id),
            exercise_name=self.ex.title,
            kind=REPORT_KIND_PROCEDURE,
        )
        self.dil_scope = ReportScope(
            exercise_id=int(self.ex.id),
            exercise_name=self.ex.title,
            kind=KIND_DILEMMA,
        )

    def tearDown(self):
        self.db.close()

    def _items(self, kind, **kwargs):
        scope = self.proc_scope if kind == REPORT_KIND_PROCEDURE else self.dil_scope
        filters = ReportFilters(**kwargs)
        return query_report_items(self.db, self.ex, scope, filters)

    def test_classification_uses_system_labels(self):
        self.assertEqual(normalize_rating_label("جيد جداً"), "جيد جدا")
        self.assertEqual(normalize_rating_label("متوسط"), "مقبول")
        self.assertEqual(list(RATING_ORDER), ["ممتاز", "جيد جدا", "جيد", "مقبول", "راسب"])

    def test_item_level_acceptable_only(self):
        items = self._items(REPORT_KIND_PROCEDURE, ratings=("مقبول",))
        self.assertEqual(len(items), 4)
        self.assertTrue(all(it.rating == "مقبول" for it in items))
        self.assertTrue(all(it.list_title == "قائمة الإنقاذ" for it in items))

    def test_each_rating_and_multiple_and_all(self):
        for rating in RATING_ORDER:
            got = self._items(REPORT_KIND_PROCEDURE, ratings=(rating,))
            self.assertTrue(all(it.rating == rating for it in got), rating)
            self.assertGreaterEqual(len(got), 1, rating)
        mixed = self._items(REPORT_KIND_PROCEDURE, ratings=("مقبول", "راسب"))
        self.assertEqual({it.rating for it in mixed}, {"مقبول", "راسب"})
        self.assertEqual(len(mixed), 5)
        all_items = self._items(REPORT_KIND_PROCEDURE, ratings=RATING_ORDER)
        self.assertEqual(len(all_items), 8)

    def test_notes_only_and_hide_notes(self):
        notes = self._items(REPORT_KIND_PROCEDURE, notes_only=True, ratings=RATING_ORDER)
        self.assertTrue(all(it.notes for it in notes))
        hidden = self._items(REPORT_KIND_PROCEDURE, show_notes=False, ratings=("ممتاز",))
        self.assertEqual(hidden[0].notes, "")

    def test_combined_filters_and_empty(self):
        items = self._items(
            REPORT_KIND_PROCEDURE,
            phase_key="preparation",
            unit_key=self.uk,
            ratings=("راسب",),
        )
        self.assertEqual(len(items), 1)
        empty = self._items(REPORT_KIND_PROCEDURE, phase_key="opening", ratings=("ممتاز",))
        self.assertEqual(empty, [])
        pcts = rating_percentages({})
        self.assertEqual(pcts["ممتاز"], 0.0)

    def test_procedure_dilemma_isolation(self):
        proc = self._items(REPORT_KIND_PROCEDURE, ratings=RATING_ORDER)
        dil = self._items(KIND_DILEMMA, ratings=RATING_ORDER)
        self.assertTrue(all("معضلة" not in it.element and it.kind == REPORT_KIND_PROCEDURE for it in proc))
        self.assertTrue(all(it.kind == KIND_DILEMMA for it in dil))
        self.assertEqual(len(dil), 2)
        proc_titles = {it.list_title for it in proc}
        dil_titles = {it.list_title for it in dil}
        self.assertNotIn("قائمة معضلة كمين", proc_titles)
        self.assertNotIn("قائمة الإنقاذ", dil_titles)

    def test_catalog_cascade_and_all_means_unrestricted(self):
        records = collect_report_lists(self.db, self.ex, self.proc_scope)
        cat = filter_catalog(records)
        self.assertTrue(any(p["key"] == "preparation" for p in cat["phases"]))
        self.assertTrue(any(u["key"] == self.uk for u in cat["units"]))
        self.assertTrue(any("أحمد المحكم" in j["label"] for j in cat["judges"]))
        self.assertTrue(any(r["list_label"] == "قائمة الإنقاذ" for r in cat["lists"]))
        all_items = self._items(REPORT_KIND_PROCEDURE)
        self.assertEqual(len(all_items), 8)

    def test_summary_counts_and_percentages_match_filter(self):
        payload = run_report(
            self.db,
            self.ex,
            self.proc_scope,
            ReportFilters(ratings=("مقبول", "راسب")),
        )
        self.assertEqual(payload["total"], 5)
        self.assertEqual(payload["counts"]["مقبول"], 4)
        self.assertEqual(payload["counts"]["راسب"], 1)
        self.assertEqual(payload["counts"]["ممتاز"], 0)
        self.assertEqual(payload["percentages"]["مقبول"], 80.0)
        self.assertEqual(payload["percentages"]["راسب"], 20.0)
        self.assertEqual(payload["unit_summary"][0]["total"], 5)
        self.assertEqual(payload["judge_summary"][0]["total"], 5)
        self.assertEqual(payload["list_summary"][0]["total"], 5)

    def test_read_only_does_not_mutate(self):
        before = self.db.query(EvaluationListSavedResult).one().payload_json
        run_report(self.db, self.ex, self.proc_scope, ReportFilters())
        after = self.db.query(EvaluationListSavedResult).one()
        self.assertEqual(after.payload_json, before)
        self.assertTrue(after.is_approved)

    def test_existing_excel_export_unchanged(self):
        saved = SimpleNamespace(
            payload_json='{"rows":[]}',
            is_approved=True,
            reopened_for_judge=False,
            total_pct=90,
            grade_label="ممتاز",
        )
        self.assertEqual(list_result_fields(saved)["position"], POSITION_APPROVED)
        raw = build_judge_results_xlsx(
            kind_label=KIND_PROCEDURE,
            phases=[
                {
                    "phase_label": "المرحلة الأولى",
                    "blocks": [
                        {
                            "unit_label": "السرية أ",
                            "judge_name": "أحمد",
                            "lists": [
                                {
                                    "title": "قائمة الإنقاذ",
                                    "pct": "80.00%",
                                    "grade": "جيد جداً",
                                    "position": POSITION_APPROVED,
                                }
                            ],
                        }
                    ],
                }
            ],
        )
        ws = load_workbook(BytesIO(raw)).active
        self.assertEqual(ws.cell(1, 1).value, KIND_PROCEDURE)
        self.assertEqual(ws.cell(5, 4).value, "جيد جداً")

    def test_excel_unified_and_summaries(self):
        payload = run_report(self.db, self.ex, self.proc_scope, ReportFilters())
        data = build_xlsx(payload)
        wb = load_workbook(BytesIO(data))
        self.assertIn("ملخص التقرير", wb.sheetnames)
        self.assertTrue(wb["ملخص التقرير"].sheet_view.rightToLeft)
        self.assertIn("ملخص النتائج حسب الوحدة", wb.sheetnames)
        self.assertIn("ملخص النتائج حسب المحكم", wb.sheetnames)
        self.assertIn("ملخص النتائج حسب قائمة التقييم", wb.sheetnames)
        self.assertIn("النتائج التفصيلية", wb.sheetnames)
        detail = wb["النتائج التفصيلية"]
        self.assertEqual(detail.cell(1, 1).value, "ت")
        self.assertEqual(detail.cell(2, 7).value, "ممتاز")
        n_rows = sum(1 for r in detail.iter_rows(min_row=2, max_col=1, values_only=True) if r[0])
        self.assertEqual(n_rows, payload["total"])
        self.assertEqual(
            report_filename(REPORT_KIND_PROCEDURE, "weakness", datetime(2026, 10, 17)),
            "تقرير_نقاط_الضعف_قوائم_الإجراءات_2026-10-17.xlsx",
        )

    def test_excel_group_by_unit_judge_list_and_safe_names(self):
        for organize in (ORGANIZE_UNIT, ORGANIZE_JUDGE, ORGANIZE_LIST):
            payload = run_report(
                self.db,
                self.ex,
                self.proc_scope,
                ReportFilters(organize=organize),
            )
            wb = load_workbook(BytesIO(build_xlsx(payload)))
            self.assertIn("ملخص التقرير", wb.sheetnames)
            self.assertGreaterEqual(len(wb.sheetnames), 5)
        used: set[str] = set()
        a = unique_sheet_name("الوحدة: أ/ب؟", used)
        b = unique_sheet_name("الوحدة: أ/ب؟", used)
        self.assertNotEqual(a, b)
        self.assertNotRegex(a, r'[\\/?*\[\]:]')
        self.assertLessEqual(len(a), 31)

    def test_preset_parser_and_long_notes(self):
        f = parse_report_filters({"report_type": "weakness"})
        self.assertEqual(tuple(f.ratings), ("مقبول", "راسب"))
        f2 = parse_report_filters({"report_type": "high_performance"})
        self.assertEqual(tuple(f2.ratings), ("ممتاز", "جيد جدا"))
        recs = collect_report_lists(self.db, self.ex, self.proc_scope)
        items = explode_scored_items(recs[0])
        long_item = [it for it in items if it.element == "بند راسب"][0]
        self.assertEqual(long_item.notes, "راسب")
        payload = run_report(self.db, self.ex, self.proc_scope, ReportFilters(ratings=("راسب",)))
        payload["items"][0].notes = "ن" * 4000
        wb = load_workbook(BytesIO(build_xlsx(payload)))
        detail = wb["النتائج التفصيلية"]
        self.assertEqual(len(str(detail.cell(2, 8).value)), 4000)


class ReportsCenterRouteTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        from app import create_app

        cls.app = create_app()
        cls.app.config["TESTING"] = True

    def _judge(self):
        user = MagicMock()
        user.role_key = RoleKey.JUDGE.value
        user.id = 1
        user.full_name = "محكم"
        user.username = "j_rep"
        return user

    def test_procedure_and_dilemma_reports_pages(self):
        user = self._judge()
        with patch("app.views.get_current_user_optional", return_value=user), patch(
            "app.views._current_workspace_exercise", return_value=None
        ):
            client = self.app.test_client()
            proc = client.get("/judge/reports/procedure")
            self.assertEqual(proc.status_code, 200)
            self.assertIn("تقارير قوائم تقييم الإجراءات".encode("utf-8"), proc.data)
            self.assertIn("التقارير".encode("utf-8"), proc.data)
            dil = client.get("/judge/reports/dilemma")
            self.assertEqual(dil.status_code, 200)
            self.assertIn("تقارير قوائم تقييم المعاضل".encode("utf-8"), dil.data)

    def test_waiting_approval_export_route_still_exists(self):
        client = self.app.test_client()
        resp = client.get("/judge/evaluation-lists/export-pending.zip")
        self.assertEqual(resp.status_code, 302)
        self.assertIn("/login", resp.headers.get("Location", ""))


if __name__ == "__main__":
    unittest.main()

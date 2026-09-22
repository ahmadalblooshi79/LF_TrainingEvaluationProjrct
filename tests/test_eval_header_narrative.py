"""صفّا وصف المعضلة ومتطلبات التنفيذ في قوائم التقييم."""
from __future__ import annotations

import io
import json
import unittest
from pathlib import Path
from tempfile import TemporaryDirectory

from openpyxl import Workbook, load_workbook

from app.evaluation_list_columns import (
    extract_eval_header_narrative,
    format_eval_narrative_cell,
    merge_eval_narrative_into_payload,
    resolve_eval_narrative_from_payload,
)
from app.evaluation_list_export import build_evaluation_list_xlsx_bytes
from app.evaluation_sheet_parser import read_evaluation_list_sheet


def _grid(*rows: list[str]) -> list[list[str]]:
    out: list[list[str]] = []
    for row in rows:
        cells = list(row)
        while len(cells) < 10:
            cells.append("")
        out.append(cells)
    return out


def _write_narrative_eval_xlsx(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "قائمة التقييم"
    ws.merge_cells("B1:J1")
    ws["B1"] = "تقييم تأمين منطقة التجمع"
    ws.merge_cells("B2:J2")
    ws["B2"] = "وصف المعضلة: عطل في آلية الذخيرة."
    ws.merge_cells("B3:J3")
    ws["B3"] = "متطلبات تنفيذ المعضلة:\n1. آلية.\n2. قافلة."
    ws["B4"] = "الوحدة:"
    ws["C4"] = "أدخل الوحدة الخاضعة للتقييم"
    ws["E4"] = "التاريخ:"
    ws["F4"] = "أدخل التاريخ"
    ws["B5"] = "قائد الوحدة:"
    ws["C5"] = "أدخل اسم قائد الوحدة"
    ws["E5"] = "المحكم:"
    ws["F5"] = "أدخل اسم المحكم"
    ws["B7"] = "عناصــــــر التقييـــــم"
    ws["E7"] = "العلامـــــــات"
    ws["I7"] = "ملاحظــــــات"
    ws["E8"] = "القصوى"
    ws["F8"] = "المكتسبة"
    ws["G8"] = "النسبة"
    ws["H8"] = "النتيجة"
    ws["B9"] = "1. بند"
    ws["E9"] = 5
    ws["F9"] = "أدخل العلامة"
    ws["B10"] = "إجمالي العلامات"
    ws["E10"] = 5
    wb.save(path)
    wb.close()


def _write_compact_eval_xlsx(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "قائمة التقييم"
    ws["B1"] = "عنوان القائمة"
    ws["B2"] = "عناصــــــر التقييـــــم"
    ws["E2"] = "العلامـــــــات"
    ws["E3"] = "القصوى"
    ws["F3"] = "المكتسبة"
    ws["B4"] = "1. بند"
    ws["E4"] = 5
    ws["F4"] = "أدخل العلامة"
    ws["B5"] = "إجمالي العلامات"
    ws["E5"] = 5
    wb.save(path)
    wb.close()


class EvalHeaderNarrativeTests(unittest.TestCase):
    def test_extract_description_and_requirements(self):
        grid = _grid(
            ["", "عنوان القائمة"],
            ["", "وصف المعضلة: يقوم السائق بالإبلاغ عن عطل."],
            ["", "متطلبات تنفيذ المعضلة:\n1. آلية.\n2. قافلة."],
            ["", "الوحدة:", "أدخل الوحدة", "", "التاريخ:"],
        )
        desc, req = extract_eval_header_narrative(grid)
        self.assertEqual(desc, "يقوم السائق بالإبلاغ عن عطل.")
        self.assertIn("آلية", req)
        self.assertIn("قافلة", req)

    def test_extract_ignores_column_headers(self):
        grid = _grid(
            ["", "عنوان القائمة"],
            ["", "عناصــــــر التقييـــــم", "", "", "العلامـــــــات"],
            ["", "", "", "", "القصوى", "المكتسبة"],
        )
        desc, req = extract_eval_header_narrative(grid)
        self.assertEqual(desc, "")
        self.assertEqual(req, "")

    def test_parser_reads_narrative_xlsx(self):
        with TemporaryDirectory() as td:
            src = Path(td) / "n.xlsx"
            _write_narrative_eval_xlsx(src)
            sheet = read_evaluation_list_sheet(src)
        self.assertEqual(sheet.get("eval_dilemma_description"), "عطل في آلية الذخيرة.")
        self.assertIn("آلية", sheet.get("eval_dilemma_requirements") or "")
        self.assertTrue(sheet.get("eval_structured"))
        self.assertGreaterEqual(len(sheet.get("eval_rows") or []), 1)

    def test_payload_saved_overrides_excel(self):
        desc, req = resolve_eval_narrative_from_payload(
            {"dilemma_description": "نص محفوظ", "dilemma_requirements": ""},
            excel_description="من الملف",
            excel_requirements="متطلب الملف",
        )
        self.assertEqual(desc, "نص محفوظ")
        self.assertEqual(req, "")

    def test_merge_keeps_header_when_tablet_omits_keys(self):
        existing = json.dumps(
            {
                "rows": [{"acquired": "4"}],
                "dilemma_description": "وصف",
                "dilemma_requirements": "متطلب",
            },
            ensure_ascii=False,
        )
        merged = merge_eval_narrative_into_payload({"rows": [{"acquired": "5"}]}, existing)
        self.assertEqual(merged["dilemma_description"], "وصف")
        self.assertEqual(merged["dilemma_requirements"], "متطلب")
        self.assertEqual(merged["rows"][0]["acquired"], "5")

    def test_export_writes_rows_2_and_3_on_narrative_template(self):
        with TemporaryDirectory() as td:
            src = Path(td) / "n.xlsx"
            _write_narrative_eval_xlsx(src)
            eval_rows = read_evaluation_list_sheet(src).get("eval_rows") or []
            data = build_evaluation_list_xlsx_bytes(
                src,
                doc_title="تقييم تأمين منطقة التجمع",
                unit_label="الكتيبة 32",
                date_str="2026-09-21",
                commander_name="قائد",
                judge_name="سالم",
                eval_rows=eval_rows,
                saved_rows=[{"acquired": "4", "notes": "", "row_kind": "score"}] * len(eval_rows),
                dilemma_description="نص محدّث للوصف",
                dilemma_requirements="1. متطلب محدّث",
            )
        wb = load_workbook(io.BytesIO(data))
        try:
            ws = wb.active
            self.assertIn("وصف المعضلة", str(ws["B2"].value or ""))
            self.assertIn("نص محدّث للوصف", str(ws["B2"].value or ""))
            self.assertIn("متطلبات تنفيذ المعضلة", str(ws["B3"].value or ""))
            self.assertIn("متطلب محدّث", str(ws["B3"].value or ""))
            self.assertEqual(ws["C4"].value, "الكتيبة 32")
            self.assertEqual(ws["F4"].value, "2026-09-21")
            self.assertEqual(ws["C5"].value, "قائد")
            self.assertEqual(ws["F5"].value, "سالم")
        finally:
            wb.close()

    def test_export_does_not_overwrite_compact_headers(self):
        with TemporaryDirectory() as td:
            src = Path(td) / "c.xlsx"
            _write_compact_eval_xlsx(src)
            eval_rows = read_evaluation_list_sheet(src).get("eval_rows") or []
            data = build_evaluation_list_xlsx_bytes(
                src,
                doc_title="عنوان القائمة",
                unit_label="وحدة",
                date_str="2026-09-21",
                commander_name="قائد",
                judge_name="سالم",
                eval_rows=eval_rows,
                saved_rows=[{"acquired": "4", "notes": "", "row_kind": "score"}] * len(eval_rows),
                dilemma_description="لا يُكتب فوق العناوين",
                dilemma_requirements="لا يُكتب فوق العناوين",
            )
        wb = load_workbook(io.BytesIO(data))
        try:
            ws = wb.active
            self.assertIn("عناص", str(ws["B2"].value or ""))
            self.assertEqual(ws["E3"].value, "القصوى")
            self.assertEqual(ws["F4"].value, 4)
        finally:
            wb.close()

    def test_format_cell_keeps_label(self):
        self.assertEqual(format_eval_narrative_cell("وصف المعضلة", ""), "وصف المعضلة:")
        self.assertTrue(
            format_eval_narrative_cell("وصف المعضلة", "نص").startswith("وصف المعضلة")
        )

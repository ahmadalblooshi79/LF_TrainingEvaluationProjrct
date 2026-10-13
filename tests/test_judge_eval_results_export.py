"""تصدير نتائج قوائم المحكمين إلى Excel."""
import unittest
from io import BytesIO
from types import SimpleNamespace

from openpyxl import load_workbook

from app.judge_eval_results_export import (
    KIND_PROCEDURE,
    POSITION_APPROVED,
    POSITION_EMPTY,
    POSITION_PENDING,
    build_judge_results_xlsx,
    list_result_fields,
)


class JudgeEvalResultsExportTests(unittest.TestCase):
    def test_position_is_one_of_the_three_labels(self):
        self.assertEqual(list_result_fields(None)["position"], POSITION_EMPTY)
        pending = SimpleNamespace(
            payload_json='{"rows":[]}',
            is_approved=False,
            reopened_for_judge=False,
            total_pct=70,
            grade_label="جيد",
        )
        self.assertEqual(list_result_fields(pending)["position"], POSITION_PENDING)
        self.assertEqual(list_result_fields(pending)["pct"], "70.00%")
        approved = SimpleNamespace(
            payload_json='{"rows":[]}',
            is_approved=True,
            reopened_for_judge=False,
            total_pct=90,
            grade_label="ممتاز",
        )
        self.assertEqual(list_result_fields(approved)["position"], POSITION_APPROVED)
        approved.reopened_for_judge = True
        self.assertEqual(list_result_fields(approved)["position"], POSITION_PENDING)

    def test_xlsx_groups_phase_then_unit_and_judge(self):
        raw = build_judge_results_xlsx(
            kind_label=KIND_PROCEDURE,
            phases=[
                {
                    "phase_label": "المرحلة الأولى",
                    "blocks": [
                        {
                            "unit_label": "السرية أ",
                            "judge_name": "أحمد",
                            "lists": [
                                {
                                    "title": "قائمة الإنقاذ",
                                    "pct": "80.00%",
                                    "grade": "جيد جداً",
                                    "position": POSITION_APPROVED,
                                },
                                {
                                    "title": "قائمة الإخلاء",
                                    "pct": "—",
                                    "grade": "—",
                                    "position": POSITION_EMPTY,
                                },
                            ],
                        }
                    ],
                }
            ],
        )
        ws = load_workbook(BytesIO(raw)).active
        self.assertTrue(ws.sheet_view.rightToLeft)
        self.assertEqual(ws.cell(1, 1).value, KIND_PROCEDURE)
        self.assertEqual(ws.cell(2, 1).value, "المرحلة: المرحلة الأولى")
        self.assertEqual(ws.cell(3, 1).value, "مسمى الوحدة: السرية أ")
        self.assertEqual(ws.cell(3, 4).value, "اسم المحكم: أحمد")
        self.assertEqual(
            [ws.cell(4, c).value for c in range(1, 6)],
            ["ت", "اسم القائمة", "النسبة المئوية", "النتيجة", "الموقف"],
        )
        self.assertEqual(ws.cell(5, 1).value, 1)
        self.assertEqual(ws.cell(5, 2).value, "قائمة الإنقاذ")
        self.assertEqual(ws.cell(5, 3).value, "80.00%")
        self.assertEqual(ws.cell(5, 4).value, "جيد جداً")
        self.assertEqual(ws.cell(5, 5).value, POSITION_APPROVED)
        self.assertEqual(ws.cell(6, 1).value, 2)
        self.assertEqual(ws.cell(6, 5).value, POSITION_EMPTY)


if __name__ == "__main__":
    unittest.main()

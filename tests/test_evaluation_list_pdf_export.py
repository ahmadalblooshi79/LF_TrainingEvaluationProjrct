import io
import unittest
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest.mock import patch

from openpyxl import Workbook, load_workbook

from app.evaluation_list_export import (
    build_evaluation_list_xlsx_bytes,
    export_download_filename,
)
from app.evaluation_sheet_parser import read_evaluation_list_sheet


def _write_sample_eval_xlsx(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "قائمة التقييم"
    ws["B1"] = "عنوان القائمة"
    ws["B2"] = "عناصــــــر التقييـــــم"
    ws["E2"] = "العلامـــــــات"
    ws["I2"] = "ملاحظــــــات"
    ws["E3"] = "القصوى"
    ws["F3"] = "المكتسبة"
    ws["G3"] = "النسبة"
    ws["H3"] = "النتيجة"
    ws["B4"] = "( أ )"
    ws["E4"] = "(ب)"
    ws["F4"] = "(جـ)"
    for i in range(1, 5):
        r = 4 + i
        ws[f"B{r}"] = f"{i}. بند"
        ws[f"E{r}"] = 5
        ws[f"F{r}"] = "أدخل العلامة"
        ws[f"G{r}"] = f'=IFERROR(F{r}/E{r},"")'
        ws[f"G{r}"].number_format = "0%"
        ws[f"H{r}"] = f'=_xlfn.IFS(G{r}>=90%,"ممتاز",G{r}>=80%,"جيد جدا",G{r}>=70%,"جيد",G{r}>=60%,"مقبول",G{r}<60%,"راسب")'
    ws["B9"] = "إجمالي العلامات"
    ws["E9"] = "=SUM(E5:E8)"
    ws["F9"] = "=SUM(F5:F8)"
    ws["B10"] = "النسبة العامة"
    ws["E10"] = "=F9/E9"
    ws["E10"].number_format = "0.00%"
    ws["B11"] = "التقدير العام"
    ws["E11"] = '=_xlfn.IFS(E10>=90%,"ممتاز",E10>=80%,"جيد جدا",E10>=70%,"جيد",E10>=60%,"مقبول",E10<60%,"راسب")'
    ws["E12"] = "المحكم:"
    ws["G12"] = "أدخل اسم المحكم"
    ws["E13"] = "التوقيع:"
    wb.save(path)
    wb.close()


class EvaluationListPdfExportTests(unittest.TestCase):
    def test_download_filename_pdf_ext(self):
        self.assertTrue(export_download_filename("قائمة.xlsx", ext="pdf").endswith(".pdf"))
        self.assertTrue(export_download_filename("قائمة", ext="pdf").endswith(".pdf"))

    def test_materialize_uses_system_scores_and_totals(self):
        with TemporaryDirectory() as td:
            src = Path(td) / "src.xlsx"
            _write_sample_eval_xlsx(src)
            eval_rows = read_evaluation_list_sheet(src).get("eval_rows") or []
            saved = [
                {"acquired": "4", "notes": "", "row_kind": "score"},
                {"acquired": "4.25", "notes": "", "row_kind": "score"},
                {"acquired": "3.5", "notes": "", "row_kind": "score"},
                {"acquired": "5", "notes": "", "row_kind": "score"},
            ]
            data = build_evaluation_list_xlsx_bytes(
                src,
                doc_title="عنوان القائمة",
                unit_label="وحدة",
                date_str="2026-09-21",
                commander_name="قائد",
                judge_name="سالم",
                eval_rows=eval_rows,
                saved_rows=saved,
                materialize_computed=True,
            )
        wb = load_workbook(io.BytesIO(data))
        try:
            ws = wb.active
            self.assertEqual(ws["F5"].value, 4)
            self.assertEqual(ws["G5"].value, 0.8)
            self.assertEqual(ws["H5"].value, "جيد جدا")
            self.assertEqual(ws["F6"].value, 4.25)
            self.assertEqual(ws["H6"].value, "جيد جدا")
            self.assertIn("BAE6FD", str(ws["H6"].fill.fgColor.rgb or "").upper())
            self.assertEqual(ws["H8"].value, "ممتاز")
            self.assertIn("BBF7D0", str(ws["H8"].fill.fgColor.rgb or "").upper())
            self.assertEqual(ws["E9"].value, 20)
            self.assertEqual(ws["F9"].value, 16.75)
            self.assertAlmostEqual(float(ws["E10"].value), 16.75 / 20.0, places=5)
            self.assertEqual(ws["E11"].value, "جيد جدا")
            self.assertEqual(ws["G12"].value, "سالم")
            self.assertFalse(str(ws["G5"].value).startswith("="))
            self.assertFalse(str(ws["E11"].value).startswith("="))
        finally:
            wb.close()

    def test_excel_export_excludes_na_from_max_and_acquired_totals(self):
        with TemporaryDirectory() as td:
            src = Path(td) / "src.xlsx"
            _write_sample_eval_xlsx(src)
            eval_rows = read_evaluation_list_sheet(src).get("eval_rows") or []
            saved = [
                {"acquired": "4", "notes": "", "row_kind": "score"},
                {"acquired": "4.25", "notes": "", "row_kind": "score"},
                {"acquired": "3.5", "notes": "", "row_kind": "score"},
                {"acquired": "na", "notes": "", "row_kind": "score"},
            ]
            data = build_evaluation_list_xlsx_bytes(
                src,
                doc_title="عنوان القائمة",
                unit_label="وحدة",
                date_str="2026-09-21",
                commander_name="قائد",
                judge_name="سالم",
                eval_rows=eval_rows,
                saved_rows=saved,
            )
        wb = load_workbook(io.BytesIO(data))
        try:
            ws = wb.active
            self.assertEqual(ws["F8"].value, "لا ينطبق")
            self.assertEqual(ws["G8"].value, "—")
            self.assertEqual(ws["H8"].value, "غير محسوب")
            self.assertEqual(ws["G5"].value, 0.8)
            self.assertEqual(ws["H5"].value, "جيد جدا")
            self.assertFalse(str(ws["G5"].value).startswith("="))
            self.assertFalse(str(ws["H5"].value).startswith("="))
            self.assertEqual(ws["E9"].value, 15)
            self.assertEqual(ws["F9"].value, 11.75)
            self.assertAlmostEqual(float(ws["E10"].value), 11.75 / 15.0, places=5)
            self.assertEqual(ws["E11"].value, "جيد")
            self.assertNotEqual(ws["E9"].value, 20)
        finally:
            wb.close()

    def test_pdf_fallback_produces_pdf(self):
        from app.evaluation_list_pdf import build_evaluation_list_pdf_bytes

        with TemporaryDirectory() as td:
            src = Path(td) / "src.xlsx"
            _write_sample_eval_xlsx(src)
            eval_rows = read_evaluation_list_sheet(src).get("eval_rows") or []
            saved = [{"acquired": "4", "notes": "ملاحظة", "row_kind": "score"} for _ in eval_rows]
            xlsx = build_evaluation_list_xlsx_bytes(
                src,
                doc_title="عنوان القائمة",
                unit_label="وحدة",
                date_str="2026-09-21",
                commander_name="قائد",
                judge_name="سالم",
                eval_rows=eval_rows,
                saved_rows=saved,
                materialize_computed=True,
            )
        with patch("app.evaluation_list_pdf.convert_xlsx_bytes_to_pdf", return_value=None):
            pdf = build_evaluation_list_pdf_bytes(xlsx)
        self.assertTrue(pdf.startswith(b"%PDF"))
        self.assertGreater(len(pdf), 500)


if __name__ == "__main__":
    unittest.main()

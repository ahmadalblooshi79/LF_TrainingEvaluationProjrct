"""ورقة Excel لتقرير التقييم التفصيلي — منفصلة عن تصدير النتائج الحالي."""
from __future__ import annotations

import io
import re

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from app.reports_center import (
    ORGANIZE_JUDGE,
    ORGANIZE_LIST,
    ORGANIZE_UNIT,
    RATING_ORDER,
    ReportItem,
)

_UNSAFE_SHEET = re.compile(r'[\\/?*\[\]:]')
_MAX_SHEET = 31
_MAX_GROUP_SHEETS = 20

_THIN = Border(
    left=Side(style="thin", color="C4B8A5"),
    right=Side(style="thin", color="C4B8A5"),
    top=Side(style="thin", color="C4B8A5"),
    bottom=Side(style="thin", color="C4B8A5"),
)
_HEAD_FILL = PatternFill("solid", fgColor="5C4033")
_META_FILL = PatternFill("solid", fgColor="F3E6D0")
_ALT_FILL = PatternFill("solid", fgColor="FBF6F0")
_WEAK_FILL = PatternFill("solid", fgColor="F8E8C8")
_FAIL_FILL = PatternFill("solid", fgColor="F8D7D3")
_HEAD_FONT = Font(name="Traditional Arabic", bold=True, color="FFFFFF", size=12)
_TITLE_FONT = Font(name="Traditional Arabic", bold=True, color="5C4033", size=14)
_BODY_FONT = Font(name="Traditional Arabic", size=12)
_BOLD = Font(name="Traditional Arabic", bold=True, size=12)
_CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
_RIGHT = Alignment(horizontal="right", vertical="center", wrap_text=True)


def sanitize_sheet_name(raw: str) -> str:
    name = _UNSAFE_SHEET.sub(" ", (raw or "").strip())
    name = re.sub(r"\s+", " ", name).strip(" '")
    if not name:
        name = "ورقة"
    return name[:_MAX_SHEET]


def unique_sheet_name(raw: str, used: set[str]) -> str:
    base = sanitize_sheet_name(raw)
    name = base
    i = 2
    while name.lower() in used:
        suffix = f"_{i}"
        name = sanitize_sheet_name(base[: _MAX_SHEET - len(suffix)] + suffix)
        i += 1
    used.add(name.lower())
    return name


def _style_sheet(ws: Worksheet) -> None:
    ws.sheet_view.rightToLeft = True
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.sheet_view.showGridLines = True


def _write_cell(ws, row, col, value, *, font=_BODY_FONT, fill=None, align=_RIGHT, number_format=None):
    cell = ws.cell(row, col, value)
    cell.font = font
    cell.alignment = align
    cell.border = _THIN
    if fill is not None:
        cell.fill = fill
    if number_format:
        cell.number_format = number_format
    return cell


def _header_row(ws, row: int, headers: tuple[str, ...]) -> None:
    for col, title in enumerate(headers, start=1):
        _write_cell(ws, row, col, title, font=_HEAD_FONT, fill=_HEAD_FILL, align=_CENTER)
    ws.row_dimensions[row].height = 22
    ws.freeze_panes = f"A{row + 1}"
    ws.auto_filter.ref = f"A{row}:{get_column_letter(len(headers))}{row}"


def _set_widths(ws: Worksheet, widths: tuple[float, ...]) -> None:
    for idx, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = width


def _rating_fill(rating: str):
    if rating == "راسب":
        return _FAIL_FILL
    if rating == "مقبول":
        return _WEAK_FILL
    return None


def _write_kv(ws, row: int, label: str, value) -> int:
    _write_cell(ws, row, 1, label, font=_BOLD, fill=_META_FILL)
    _write_cell(ws, row, 2, value if value not in (None, "") else "—", font=_BODY_FONT, fill=_META_FILL)
    ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=4)
    return row + 1


def _extend_autofilter(ws: Worksheet, header_row: int, last_row: int, cols: int) -> None:
    if last_row < header_row:
        last_row = header_row
    ws.auto_filter.ref = f"A{header_row}:{get_column_letter(cols)}{last_row}"


def _write_summary_sheet(ws: Worksheet, payload: dict) -> None:
    _style_sheet(ws)
    ws.cell(1, 1, "ملخص التقرير").font = _TITLE_FONT
    ws.merge_cells("A1:D1")
    ws.row_dimensions[1].height = 26
    row = 3
    notes_opt = "إظهار الملاحظات"
    if payload.get("notes_only"):
        notes_opt += " — العناصر ذات الملاحظات فقط"
    meta = [
        ("اسم التمرين", payload.get("exercise_name") or "—"),
        ("نوع قوائم التقييم", payload.get("kind_label") or "—"),
        ("نوع التقرير", payload.get("report_type_label") or "—"),
        ("المرحلة", payload.get("phase_label") or "جميع المراحل"),
        ("الوحدة", payload.get("unit_label") or "جميع الوحدات"),
        ("المحكم", payload.get("judge_label") or "جميع المحكمين"),
        ("قائمة التقييم", payload.get("list_label") or "جميع القوائم"),
        ("نتائج التقييم المختارة", "، ".join(payload.get("ratings_selected") or []) or "—"),
        ("خيار الملاحظات", notes_opt),
        ("طريقة تنظيم التقرير", payload.get("organize_label") or "—"),
        ("تاريخ ووقت إنشاء التقرير", payload.get("generated_at") or "—"),
        ("نطاق البيانات", payload.get("scope_note") or "—"),
        ("إجمالي العناصر المطابقة", int(payload.get("total") or 0)),
    ]
    for label, value in meta:
        row = _write_kv(ws, row, label, value)
    row += 1
    _write_cell(ws, row, 1, "النتيجة", font=_HEAD_FONT, fill=_HEAD_FILL, align=_CENTER)
    _write_cell(ws, row, 2, "العدد", font=_HEAD_FONT, fill=_HEAD_FILL, align=_CENTER)
    _write_cell(ws, row, 3, "النسبة", font=_HEAD_FONT, fill=_HEAD_FILL, align=_CENTER)
    counts = payload.get("counts") or {}
    pcts = payload.get("percentages") or {}
    row += 1
    for rating in RATING_ORDER:
        fill = _rating_fill(rating)
        _write_cell(ws, row, 1, rating, fill=fill, align=_CENTER)
        _write_cell(ws, row, 2, int(counts.get(rating) or 0), align=_CENTER, number_format="0")
        _write_cell(
            ws,
            row,
            3,
            float(pcts.get(rating) or 0) / 100.0,
            align=_CENTER,
            number_format="0.0%",
            fill=fill,
        )
        row += 1
    _set_widths(ws, (36, 28, 18, 42))
    ws.freeze_panes = "A3"


def _write_unit_summary_sheet(ws: Worksheet, rows: list[dict]) -> None:
    _style_sheet(ws)
    headers = (
        "الوحدة",
        "إجمالي العناصر",
        "ممتاز",
        "جيد جدا",
        "جيد",
        "مقبول",
        "راسب",
        "نسبة ممتاز",
        "نسبة جيد جدا",
        "نسبة جيد",
        "نسبة مقبول",
        "نسبة راسب",
    )
    _header_row(ws, 1, headers)
    r = 2
    for item in rows:
        counts = item.get("counts") or {}
        pcts = item.get("pcts") or {}
        vals = [
            item.get("unit_label") or "—",
            int(item.get("total") or 0),
        ]
        for rating in RATING_ORDER:
            vals.append(int(counts.get(rating) or 0))
        for rating in RATING_ORDER:
            vals.append(float(pcts.get(rating) or 0) / 100.0)
        for col, val in enumerate(vals, start=1):
            fmt = "0.0%" if col >= 8 else ("0" if col >= 2 else None)
            align = _RIGHT if col == 1 else _CENTER
            _write_cell(ws, r, col, val, align=align, number_format=fmt)
        r += 1
    if r == 2:
        _write_cell(ws, r, 1, "لا توجد نتائج مطابقة للفلاتر المحددة")
        r += 1
    _extend_autofilter(ws, 1, r - 1, len(headers))
    _set_widths(ws, (28, 16, 12, 12, 12, 12, 12, 14, 14, 14, 14, 14))


def _write_judge_summary_sheet(ws: Worksheet, rows: list[dict]) -> None:
    _style_sheet(ws)
    headers = (
        "المحكم",
        "الوحدة",
        "إجمالي العناصر",
        "ممتاز",
        "جيد جدا",
        "جيد",
        "مقبول",
        "راسب",
    )
    _header_row(ws, 1, headers)
    r = 2
    for item in rows:
        counts = item.get("counts") or {}
        vals = [
            item.get("judge_label") or "—",
            item.get("unit_label") or "—",
            int(item.get("total") or 0),
        ]
        for rating in RATING_ORDER:
            vals.append(int(counts.get(rating) or 0))
        for col, val in enumerate(vals, start=1):
            fmt = "0" if col >= 3 else None
            align = _CENTER if col >= 3 else _RIGHT
            _write_cell(ws, r, col, val, align=align, number_format=fmt)
        r += 1
    if r == 2:
        _write_cell(ws, r, 1, "لا توجد نتائج مطابقة للفلاتر المحددة")
        r += 1
    _extend_autofilter(ws, 1, r - 1, len(headers))
    _set_widths(ws, (28, 28, 16, 12, 12, 12, 12, 12))


def _write_list_summary_sheet(ws: Worksheet, rows: list[dict]) -> None:
    _style_sheet(ws)
    headers = (
        "اسم القائمة",
        "الوحدة",
        "المحكم",
        "إجمالي العناصر",
        "ممتاز",
        "جيد جدا",
        "جيد",
        "مقبول",
        "راسب",
    )
    _header_row(ws, 1, headers)
    r = 2
    for item in rows:
        counts = item.get("counts") or {}
        vals = [
            item.get("list_title") or "—",
            item.get("unit_label") or "—",
            item.get("judge_label") or "—",
            int(item.get("total") or 0),
        ]
        for rating in RATING_ORDER:
            vals.append(int(counts.get(rating) or 0))
        for col, val in enumerate(vals, start=1):
            fmt = "0" if col >= 4 else None
            align = _CENTER if col >= 4 else _RIGHT
            _write_cell(ws, r, col, val, align=align, number_format=fmt)
        ws.row_dimensions[r].height = 22
        r += 1
    if r == 2:
        _write_cell(ws, r, 1, "لا توجد نتائج مطابقة للفلاتر المحددة")
        r += 1
    _extend_autofilter(ws, 1, r - 1, len(headers))
    _set_widths(ws, (40, 24, 24, 16, 12, 12, 12, 12, 12))


def _write_detail_rows(
    ws: Worksheet,
    items: list[ReportItem],
    *,
    headers: tuple[str, ...],
    row_values,
    widths: tuple[float, ...],
) -> None:
    _style_sheet(ws)
    _header_row(ws, 1, headers)
    r = 2
    for seq, item in enumerate(items, start=1):
        values = row_values(seq, item)
        alt = _ALT_FILL if seq % 2 == 0 else None
        rating = item.rating
        for col, val in enumerate(values, start=1):
            fill = _rating_fill(rating) if headers[col - 1] == "النتيجة" else alt
            align = _CENTER if col == 1 or headers[col - 1] in {"النتيجة"} else _RIGHT
            _write_cell(ws, r, col, val, fill=fill, align=align)
        ws.row_dimensions[r].height = 24
        r += 1
    if r == 2:
        _write_cell(ws, r, 1, "لا توجد نتائج مطابقة للفلاتر المحددة")
        r += 1
    _extend_autofilter(ws, 1, r - 1, len(headers))
    _set_widths(ws, widths)


def _unified_values(seq: int, item: ReportItem) -> tuple:
    return (
        seq,
        item.phase_label,
        item.unit_label,
        item.judge_label,
        item.list_title,
        item.element,
        item.rating,
        item.notes,
    )


_UNIFIED_HEADERS = (
    "ت",
    "المرحلة",
    "الوحدة",
    "المحكم",
    "اسم قائمة التقييم",
    "عنصر التقييم",
    "النتيجة",
    "الملاحظات",
)
_UNIFIED_WIDTHS = (8, 22, 22, 22, 36, 44, 14, 40)


def _group_items(items: list[ReportItem], key_fn) -> list[tuple[str, list[ReportItem]]]:
    buckets: dict[str, list[ReportItem]] = {}
    order: list[str] = []
    for it in items:
        key = key_fn(it) or "—"
        if key not in buckets:
            buckets[key] = []
            order.append(key)
        buckets[key].append(it)
    return [(k, buckets[k]) for k in order]


def _write_grouped_sheets(
    wb: Workbook,
    used: set[str],
    groups: list[tuple[str, list[ReportItem]]],
    *,
    headers: tuple[str, ...],
    row_values,
    widths: tuple[float, ...],
    fallback_title: str,
) -> None:
    if len(groups) > _MAX_GROUP_SHEETS:
        ws = wb.create_sheet(unique_sheet_name(fallback_title, used))
        _style_sheet(ws)
        row = 1
        seq = 1
        for label, group in groups:
            _write_cell(ws, row, 1, label, font=_HEAD_FONT, fill=_HEAD_FILL)
            ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=len(headers))
            row += 1
            for col, title in enumerate(headers, start=1):
                _write_cell(ws, row, col, title, font=_HEAD_FONT, fill=_HEAD_FILL, align=_CENTER)
            row += 1
            for item in group:
                values = row_values(seq, item)
                for col, val in enumerate(values, start=1):
                    fill = _rating_fill(item.rating) if headers[col - 1] == "النتيجة" else None
                    _write_cell(ws, row, col, val, fill=fill, align=_CENTER if col == 1 else _RIGHT)
                ws.row_dimensions[row].height = 24
                row += 1
                seq += 1
            row += 1
        _set_widths(ws, widths)
        return
    for label, group in groups:
        ws = wb.create_sheet(unique_sheet_name(label, used))
        _write_detail_rows(ws, group, headers=headers, row_values=row_values, widths=widths)


def build_detailed_report_xlsx(payload: dict) -> bytes:
    items: list[ReportItem] = list(payload.get("items") or [])
    organize = payload.get("organize") or ORGANIZE_UNIT
    wb = Workbook()
    used: set[str] = set()
    ws0 = wb.active
    ws0.title = unique_sheet_name("ملخص التقرير", used)
    _write_summary_sheet(ws0, payload)

    ws_u = wb.create_sheet(unique_sheet_name("ملخص النتائج حسب الوحدة", used))
    _write_unit_summary_sheet(ws_u, payload.get("unit_summary") or [])
    ws_j = wb.create_sheet(unique_sheet_name("ملخص النتائج حسب المحكم", used))
    _write_judge_summary_sheet(ws_j, payload.get("judge_summary") or [])
    ws_l = wb.create_sheet(unique_sheet_name("ملخص النتائج حسب قائمة التقييم", used))
    _write_list_summary_sheet(ws_l, payload.get("list_summary") or [])

    if organize == ORGANIZE_UNIT:
        groups = _group_items(items, lambda it: it.unit_label)
        _write_grouped_sheets(
            wb,
            used,
            groups,
            headers=("ت", "المحكم", "اسم قائمة التقييم", "عنصر التقييم", "النتيجة", "الملاحظات"),
            row_values=lambda seq, it: (
                seq,
                it.judge_label,
                it.list_title,
                it.element,
                it.rating,
                it.notes,
            ),
            widths=(8, 22, 36, 44, 14, 40),
            fallback_title="النتائج حسب الوحدة",
        )
    elif organize == ORGANIZE_JUDGE:
        groups = _group_items(items, lambda it: it.judge_label)
        _write_grouped_sheets(
            wb,
            used,
            groups,
            headers=("ت", "المرحلة", "الوحدة", "اسم قائمة التقييم", "عنصر التقييم", "النتيجة", "الملاحظات"),
            row_values=lambda seq, it: (
                seq,
                it.phase_label,
                it.unit_label,
                it.list_title,
                it.element,
                it.rating,
                it.notes,
            ),
            widths=(8, 22, 22, 36, 44, 14, 40),
            fallback_title="النتائج حسب المحكم",
        )
    elif organize == ORGANIZE_LIST:
        groups = _group_items(items, lambda it: it.list_title)
        _write_grouped_sheets(
            wb,
            used,
            groups,
            headers=("ت", "المرحلة", "الوحدة", "المحكم", "عنصر التقييم", "النتيجة", "الملاحظات"),
            row_values=lambda seq, it: (
                seq,
                it.phase_label,
                it.unit_label,
                it.judge_label,
                it.element,
                it.rating,
                it.notes,
            ),
            widths=(8, 22, 22, 22, 44, 14, 40),
            fallback_title="النتائج حسب القائمة",
        )
    else:
        ws = wb.create_sheet(unique_sheet_name("النتائج التفصيلية", used))
        _write_detail_rows(
            ws,
            items,
            headers=_UNIFIED_HEADERS,
            row_values=_unified_values,
            widths=_UNIFIED_WIDTHS,
        )

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()

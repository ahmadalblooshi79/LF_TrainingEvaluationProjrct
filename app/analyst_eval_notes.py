# -*- coding: utf-8 -*-
"""ملاحظات قوائم التقييم وقوائم تقييم المعاضل — عرض وتصدير لمساحة المحللين."""
from __future__ import annotations

import io
import json
import re

from sqlalchemy.orm import Session

from app.models import (
    EvaluationListPdfItem,
    EvaluationListSavedResult,
    Exercise,
    ExercisePlannerFlowBundle,
    ExercisePlannerFlowBundleActionEval,
    PlannerFlowBundleEvalSavedResult,
)
from app.unit_levels_catalog import label_for_unit_level_key

SOURCE_EVAL = "eval"
SOURCE_ACTION = "action"
SOURCE_ALL = "all"

SOURCE_LABELS = {
    SOURCE_EVAL: "قوائم التقييم",
    SOURCE_ACTION: "قوائم تقييم المعاضل",
}

_UNSAFE_NAME_RE = re.compile(r'[<>:"/\\|?*\x00-\x1f]')


def _parse_payload_rows(payload_json: str | None) -> list[dict]:
    if not (payload_json or "").strip():
        return []
    try:
        data = json.loads(payload_json)
    except Exception:
        return []
    if not isinstance(data, dict):
        return []
    rows = data.get("rows") or []
    return rows if isinstance(rows, list) else []


def _row_element_label(row: dict) -> str:
    for key in ("element", "text", "criterion", "label", "title"):
        val = str(row.get(key) or "").strip()
        if val:
            return val[:400]
    return ""


def iter_payload_notes(payload_json: str | None) -> list[dict[str, str]]:
    """صفوف التقييم ذات ملاحظة غير فارغة."""
    out: list[dict[str, str]] = []
    for idx, row in enumerate(_parse_payload_rows(payload_json)):
        if not isinstance(row, dict):
            continue
        if str(row.get("row_kind") or "score").strip().lower() == "section":
            continue
        note = str(row.get("notes") or "").strip()
        if not note:
            continue
        out.append(
            {
                "element": _row_element_label(row) or f"بند {idx + 1}",
                "note": note[:4000],
            }
        )
    return out


def _unit_order_index(unit_key: str) -> int:
    from app.unit_levels_catalog import UNIT_LEVELS

    uk = (unit_key or "").strip()
    for idx, row in enumerate(UNIT_LEVELS):
        if (row.get("key") or "") == uk:
            return idx
    return len(UNIT_LEVELS) + 1


def _unit_label(db: Session, unit_key: str, fallback: str = "") -> str:
    uk = (unit_key or "").strip()
    return (
        (fallback or "").strip()
        or label_for_unit_level_key(uk, db=db)
        or uk
        or "—"
    )


def _canonical_eval_saved_map(
    db: Session, exercise_id: int, item_ids: list[int]
) -> dict[int, EvaluationListSavedResult]:
    if not item_ids:
        return {}
    rows = (
        db.query(EvaluationListSavedResult)
        .filter(EvaluationListSavedResult.evaluation_item_id.in_(item_ids))
        .order_by(
            EvaluationListSavedResult.updated_at.desc(),
            EvaluationListSavedResult.id.desc(),
        )
        .all()
    )
    out: dict[int, EvaluationListSavedResult] = {}
    for r in rows:
        iid = int(r.evaluation_item_id)
        prev = out.get(iid)
        if prev is None:
            out[iid] = r
            continue
        rid = int(getattr(r, "exercise_id", 0) or 0)
        pid = int(getattr(prev, "exercise_id", 0) or 0)
        if rid == int(exercise_id) and pid != int(exercise_id):
            out[iid] = r
    return out


def _empty_unit_bucket(unit_key: str, unit_label: str) -> dict:
    return {
        "unit_key": unit_key,
        "unit_label": unit_label,
        "eval_notes": [],
        "action_notes": [],
    }


def build_entered_eval_notes_report(db: Session, exercise: Exercise) -> dict:
    """كل الملاحظات المدخلة — حسب تسلسل التنظيم، مفصولة حسب نوع القائمة."""
    by_unit: dict[str, dict] = {}

    items = (
        db.query(EvaluationListPdfItem)
        .filter(EvaluationListPdfItem.exercise_id == int(exercise.id))
        .all()
    )
    item_ids = [int(it.id) for it in items if getattr(it, "id", None) is not None]
    saved_by_item: dict[int, EvaluationListSavedResult] = {}
    if item_ids:
        saved_by_item = _canonical_eval_saved_map(db, int(exercise.id), item_ids)
    for it in items:
        saved = saved_by_item.get(int(it.id))
        if saved is None:
            continue
        notes = iter_payload_notes(getattr(saved, "payload_json", None))
        if not notes:
            continue
        uk = (it.unit_level_key or "").strip()
        bucket = by_unit.setdefault(
            uk,
            _empty_unit_bucket(
                uk,
                _unit_label(db, uk, getattr(it, "unit_level_label", "") or ""),
            ),
        )
        list_title = (getattr(it, "text", None) or "قائمة تقييم").strip()
        for n in notes:
            bucket["eval_notes"].append(
                {
                    "source": SOURCE_EVAL,
                    "source_label": SOURCE_LABELS[SOURCE_EVAL],
                    "list_title": list_title,
                    "element": n["element"],
                    "note": n["note"],
                }
            )

    saved_action = (
        db.query(PlannerFlowBundleEvalSavedResult, ExercisePlannerFlowBundleActionEval)
        .join(
            ExercisePlannerFlowBundleActionEval,
            ExercisePlannerFlowBundleActionEval.id
            == PlannerFlowBundleEvalSavedResult.bundle_action_eval_id,
        )
        .filter(PlannerFlowBundleEvalSavedResult.exercise_id == int(exercise.id))
        .all()
    )
    for saved, slot in saved_action:
        notes = iter_payload_notes(getattr(saved, "payload_json", None))
        if not notes:
            continue
        uk = (saved.unit_level_key or "").strip()
        if not uk:
            bundle = db.get(ExercisePlannerFlowBundle, int(slot.bundle_id))
            uk = (bundle.unit_level_key or "").strip() if bundle is not None else ""
        bucket = by_unit.setdefault(
            uk,
            _empty_unit_bucket(uk, _unit_label(db, uk)),
        )
        list_title = (slot.title or "قائمة تقييم إجراءات").strip()
        for n in notes:
            bucket["action_notes"].append(
                {
                    "source": SOURCE_ACTION,
                    "source_label": SOURCE_LABELS[SOURCE_ACTION],
                    "list_title": list_title,
                    "element": n["element"],
                    "note": n["note"],
                }
            )

    units = sorted(
        by_unit.values(),
        key=lambda u: (_unit_order_index(u["unit_key"]), u["unit_label"], u["unit_key"]),
    )
    n_eval = sum(len(u["eval_notes"]) for u in units)
    n_action = sum(len(u["action_notes"]) for u in units)
    return {
        "has_exercise": True,
        "units": units,
        "eval_count": n_eval,
        "action_count": n_action,
        "notes_count": n_eval + n_action,
        "units_count": len(units),
    }


def flatten_note_rows(report: dict, *, source: str) -> list[dict]:
    want = (source or SOURCE_ALL).strip().lower()
    if want not in (SOURCE_EVAL, SOURCE_ACTION, SOURCE_ALL):
        want = SOURCE_ALL
    rows: list[dict] = []
    seq = 0
    for unit in report.get("units") or []:
        chunks: list[list[dict]] = []
        if want in (SOURCE_EVAL, SOURCE_ALL):
            chunks.append(list(unit.get("eval_notes") or []))
        if want in (SOURCE_ACTION, SOURCE_ALL):
            chunks.append(list(unit.get("action_notes") or []))
        for group in chunks:
            for item in group:
                seq += 1
                rows.append(
                    {
                        "seq": seq,
                        "unit_key": unit.get("unit_key") or "",
                        "unit_label": unit.get("unit_label") or "—",
                        "source": item.get("source") or "",
                        "source_label": item.get("source_label") or "",
                        "list_title": item.get("list_title") or "",
                        "element": item.get("element") or "",
                        "note": item.get("note") or "",
                    }
                )
    return rows


def _export_basename(source: str) -> str:
    if source == SOURCE_EVAL:
        return "ملاحظات_قوائم_التقييم"
    if source == SOURCE_ACTION:
        return "ملاحظات_قوائم_تقييم_المعاضل"
    return "جميع_الملاحظات"


def notes_export_filename(source: str, ext: str) -> str:
    name = _UNSAFE_NAME_RE.sub("_", _export_basename(source)).strip(" .")
    ext = (ext or "").lstrip(".")
    return f"{name}.{ext}"


def build_notes_xlsx_bytes(rows: list[dict], *, title: str) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "الملاحظات"
    ws.sheet_view.rightToLeft = True
    headers = ("ت", "الوحدة", "نوع القائمة", "مسمى التقييم", "البند", "الملاحظة")
    head_fill = PatternFill("solid", fgColor="5C4033")
    head_font = Font(name="Traditional Arabic", bold=True, color="FFFFFF", size=12)
    body_font = Font(name="Traditional Arabic", size=12)
    thin = Border(
        left=Side(style="thin", color="C4B8A5"),
        right=Side(style="thin", color="C4B8A5"),
        top=Side(style="thin", color="C4B8A5"),
        bottom=Side(style="thin", color="C4B8A5"),
    )
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=6)
    title_cell = ws.cell(1, 1, title)
    title_cell.font = Font(name="Traditional Arabic", bold=True, size=14, color="5C4033")
    title_cell.alignment = Alignment(horizontal="right", vertical="center")
    for col, h in enumerate(headers, start=1):
        cell = ws.cell(2, col, h)
        cell.fill = head_fill
        cell.font = head_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin
    for i, row in enumerate(rows, start=3):
        values = (
            row.get("seq"),
            row.get("unit_label"),
            row.get("source_label"),
            row.get("list_title"),
            row.get("element"),
            row.get("note"),
        )
        fill = PatternFill("solid", fgColor="FBF6F0") if i % 2 == 0 else None
        for col, val in enumerate(values, start=1):
            cell = ws.cell(i, col, val if val is not None else "")
            cell.font = body_font
            cell.alignment = Alignment(
                horizontal="right" if col > 1 else "center",
                vertical="top",
                wrap_text=True,
            )
            cell.border = thin
            if fill is not None:
                cell.fill = fill
    widths = (6, 28, 24, 36, 36, 48)
    for idx, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = w
    ws.freeze_panes = "A3"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def build_notes_docx_bytes(rows: list[dict], *, title: str) -> bytes:
    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor

    doc = Document()
    section = doc.sections[0]
    sect_pr = section._sectPr
    if sect_pr.find(qn("w:bidi")) is None:
        sect_pr.append(OxmlElement("w:bidi"))
    section.page_width = Cm(29.7)
    section.page_height = Cm(21.0)
    section.right_margin = Cm(1.4)
    section.left_margin = Cm(1.4)
    section.top_margin = Cm(1.4)
    section.bottom_margin = Cm(1.4)

    def _rtl_run(run, *, size=12, bold=False, color=None):
        run.bold = bold
        run.font.size = Pt(size)
        run.font.name = "Traditional Arabic"
        if color is not None:
            run.font.color.rgb = color
        rpr = run._element.get_or_add_rPr()
        rfonts = rpr.find(qn("w:rFonts"))
        if rfonts is None:
            rfonts = OxmlElement("w:rFonts")
            rpr.insert(0, rfonts)
        rfonts.set(qn("w:ascii"), "Traditional Arabic")
        rfonts.set(qn("w:hAnsi"), "Traditional Arabic")
        rfonts.set(qn("w:cs"), "Traditional Arabic")
        if rpr.find(qn("w:rtl")) is None:
            rpr.append(OxmlElement("w:rtl"))

    def _rtl_para(p, align=WD_ALIGN_PARAGRAPH.RIGHT):
        p.alignment = align
        ppr = p._element.get_or_add_pPr()
        if ppr.find(qn("w:bidi")) is None:
            ppr.append(OxmlElement("w:bidi"))

    heading = doc.add_paragraph()
    _rtl_para(heading)
    run = heading.add_run(title)
    _rtl_run(run, size=18, bold=True, color=RGBColor(92, 64, 51))

    headers = ("ت", "الوحدة", "نوع القائمة", "مسمى التقييم", "البند", "الملاحظة")
    table = doc.add_table(rows=1 + max(len(rows), 1), cols=6)
    table.alignment = WD_TABLE_ALIGNMENT.RIGHT
    table.style = "Table Grid"
    for col, text in enumerate(headers):
        cell = table.rows[0].cells[col]
        cell.text = ""
        p = cell.paragraphs[0]
        _rtl_para(p, WD_ALIGN_PARAGRAPH.CENTER)
        r = p.add_run(text)
        _rtl_run(r, size=11, bold=True, color=RGBColor(255, 255, 255))
        shading = OxmlElement("w:shd")
        shading.set(qn("w:fill"), "5C4033")
        shading.set(qn("w:val"), "clear")
        cell._tePr = cell._tc.get_or_add_tcPr()
        cell._tc.get_or_add_tcPr().append(shading)
    if not rows:
        cell = table.rows[1].cells[0]
        cell.merge(table.rows[1].cells[5])
        p = cell.paragraphs[0]
        _rtl_para(p)
        r = p.add_run("لا توجد ملاحظات مدخلة.")
        _rtl_run(r, size=12)
    else:
        for i, row in enumerate(rows):
            values = (
                str(row.get("seq") or ""),
                str(row.get("unit_label") or ""),
                str(row.get("source_label") or ""),
                str(row.get("list_title") or ""),
                str(row.get("element") or ""),
                str(row.get("note") or ""),
            )
            for col, text in enumerate(values):
                cell = table.rows[i + 1].cells[col]
                cell.text = ""
                p = cell.paragraphs[0]
                _rtl_para(p)
                r = p.add_run(text)
                _rtl_run(r, size=11)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()

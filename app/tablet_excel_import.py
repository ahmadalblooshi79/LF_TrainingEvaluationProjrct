"""استيراد نتائج التابلت من Excel/ZIP مع معاينة ومطابقة متعددة المفاتيح."""
from __future__ import annotations

import hashlib
import json
import shutil
import zipfile
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from sqlalchemy.orm import Session

from app.eval_identity_matcher import (
    CandidateTarget,
    MatchResult,
    SourceIdentity,
    load_candidates,
    match_evaluation,
    remember_mapping,
    record_transfer_audit,
)
from app.paths import data_dir
from app.tablet_recovery import (
    _ensure_iterable,
    _write_saved_result,
    parse_pending_body,
)

EXPORT_FORMAT_VERSION = "1"
METADATA_SHEET = "_System_Metadata"


def _uploads_root() -> Path:
    root = data_dir() / "tablet_excel_imports"
    root.mkdir(parents=True, exist_ok=True)
    return root


def sha256_bytes(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def sha256_file(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def _as_int(v: Any) -> int | None:
    if v is None or v == "":
        return None
    try:
        n = int(str(v).strip())
    except (TypeError, ValueError):
        return None
    return n if n > 0 else None


def _as_str(v: Any) -> str:
    if v is None:
        return ""
    return str(v).strip()


def parse_metadata_sheet(ws) -> dict[str, str]:
    out: dict[str, str] = {}
    for row in ws.iter_rows(min_row=1, values_only=True):
        if not row:
            continue
        key = _as_str(row[0] if len(row) > 0 else "")
        val = _as_str(row[1] if len(row) > 1 else "")
        if key:
            out[key] = val
    return out


def parse_results_sheet(ws) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    headers: list[str] = []
    for i, row in enumerate(ws.iter_rows(min_row=1, values_only=True)):
        vals = ["" if c is None else str(c).strip() for c in row]
        if i == 0:
            headers = [h or f"c{j}" for j, h in enumerate(vals)]
            continue
        if not any(vals):
            continue
        rec = {headers[j]: vals[j] if j < len(vals) else "" for j in range(len(headers))}
        rows.append(
            {
                "row_kind": rec.get("row_kind") or rec.get("نوع الصف") or "criterion",
                "element": rec.get("element") or rec.get("العنصر") or rec.get("عناصر التقييم") or "",
                "max_val": rec.get("max_val") or rec.get("القصوى") or rec.get("العلامة القصوى") or "",
                "acquired": rec.get("acquired") or rec.get("المكتسبة") or rec.get("العلامة المكتسبة") or "",
                "notes": rec.get("notes") or rec.get("الملاحظات") or rec.get("ملاحظات") or "",
            }
        )
    return rows


def _read_xlsx_bytes(data: bytes) -> tuple[dict[str, str], list[dict[str, Any]], str]:
    from io import BytesIO

    from openpyxl import load_workbook

    wb = load_workbook(BytesIO(data), read_only=True, data_only=True)
    meta: dict[str, str] = {}
    result_rows: list[dict[str, Any]] = []
    list_name = ""
    try:
        names = wb.sheetnames
        if METADATA_SHEET in names:
            meta = parse_metadata_sheet(wb[METADATA_SHEET])
        for name in names:
            if name == METADATA_SHEET:
                continue
            if not list_name:
                list_name = name
            parsed = parse_results_sheet(wb[name])
            if parsed:
                result_rows = parsed
                list_name = name
                break
    finally:
        wb.close()
    if not meta.get("evaluation_list_name") and list_name:
        meta["evaluation_list_name"] = list_name
    return meta, result_rows, list_name


def _source_from_meta(meta: dict[str, str], rows: list[dict[str, Any]]) -> SourceIdentity:
    return SourceIdentity(
        exercise_id=_as_int(meta.get("exercise_id")),
        exercise_name=_as_str(meta.get("exercise_name")),
        unit_id=_as_str(meta.get("unit_id")),
        unit_name=_as_str(meta.get("unit_name")),
        evaluation_list_id=_as_int(meta.get("evaluation_list_id")),
        evaluation_list_name=_as_str(meta.get("evaluation_list_name")),
        eval_item_id=_as_int(meta.get("eval_item_id")),
        evaluation_type=_as_str(meta.get("evaluation_type") or "evaluation_list"),
        stage_id=_as_str(meta.get("stage_id")),
        stage_name=_as_str(meta.get("stage_name")),
        dilemma_id=_as_int(meta.get("dilemma_id")),
        dilemma_name=_as_str(meta.get("dilemma_name") or meta.get("dilemma_description")),
        judge_id=_as_int(meta.get("judge_id")),
        judge_name=_as_str(meta.get("judge_name")),
        client_op_id=_as_str(meta.get("client_op_id")),
        device_id=_as_str(meta.get("device_id")),
    )


@dataclass
class ImportItemPreview:
    source: SourceIdentity
    rows: list[dict[str, Any]]
    payload: dict[str, Any]
    match: MatchResult
    approved_source: bool = False
    saved_at: str = ""
    approved_at: str = ""
    media_count: int = 0
    skip_reason: str = ""
    duplicate: bool = False

    def to_dict(self) -> dict[str, Any]:
        src = self.source
        tgt = self.match.target
        return {
            "source_eval_item_id": src.eval_item_id,
            "source_unit_id": src.unit_id,
            "source_unit_name": src.unit_name,
            "source_list_name": src.evaluation_list_name,
            "source_judge_id": src.judge_id,
            "source_judge_name": src.judge_name,
            "source_exercise_id": src.exercise_id,
            "client_op_id": src.client_op_id,
            "target_eval_item_id": tgt.eval_item_id if tgt else None,
            "target_unit_id": tgt.unit_id if tgt else "",
            "target_unit_name": tgt.unit_name if tgt else "",
            "target_list_name": tgt.list_name if tgt else "",
            "target_judge_id": tgt.judge_id if tgt else None,
            "target_judge_name": tgt.judge_name if tgt else "",
            "match_method": self.match.match_method,
            "match_status": self.match.status,
            "match_status_ar": self.match.status_ar,
            "conflicts": self.match.conflicts,
            "candidates": [
                {
                    "eval_item_id": c.eval_item_id,
                    "unit_name": c.unit_name,
                    "list_name": c.list_name,
                }
                for c in self.match.candidates
            ],
            "approved_source": self.approved_source,
            "saved_at": self.saved_at,
            "approved_at": self.approved_at,
            "row_count": len(self.rows),
            "media_count": self.media_count,
            "skip_reason": self.skip_reason,
            "duplicate": self.duplicate,
            "ready": self.match.status
            in ("MATCHED_EXACT", "MATCHED_BY_UNIT_LIST", "MATCHED_BY_CONTEXT")
            and not self.duplicate
            and self.match.status != "CONFLICT",
        }


@dataclass
class ImportPreview:
    package_hash: str
    package_name: str
    extract_dir: str
    previously_applied: dict[str, Any] | None
    items: list[ImportItemPreview] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)
    media_manifest: list[dict[str, Any]] = field(default_factory=list)

    def to_session(self) -> dict[str, Any]:
        return {
            "package_hash": self.package_hash,
            "package_name": self.package_name,
            "extract_dir": self.extract_dir,
        }


def _client_op_seen(db: Session, client_op_id: str) -> bool:
    if not client_op_id:
        return False
    from app.models import TabletClientOp
    from app.models.tablet_transfer import TabletTransferAudit

    if (
        db.query(TabletClientOp)
        .filter(TabletClientOp.client_op_id == client_op_id)
        .first()
        is not None
    ):
        return True
    return (
        db.query(TabletTransferAudit)
        .filter(TabletTransferAudit.client_op_id == client_op_id)
        .first()
        is not None
    )


def _package_applied(db: Session, package_hash: str):
    from app.models.tablet_transfer import TabletImportPackage

    return (
        db.query(TabletImportPackage)
        .filter(TabletImportPackage.package_hash == package_hash)
        .one_or_none()
    )


def extract_import_archive(src: Path, dest: Path) -> None:
    dest.mkdir(parents=True, exist_ok=True)
    if src.suffix.lower() == ".xlsx":
        shutil.copy2(src, dest / src.name)
        return
    with zipfile.ZipFile(src, "r") as zf:
        zf.extractall(dest)


def _find_xlsx(extract_dir: Path) -> Path | None:
    files = list(extract_dir.rglob("*.xlsx"))
    if not files:
        return None
    preferred = [p for p in files if not p.name.startswith("~$")]
    files = preferred or files
    meta_first = [p for p in files if "Evaluation_" in p.name or "evaluation" in p.name.lower()]
    return (meta_first or files)[0]


def _load_manifest(extract_dir: Path) -> dict[str, Any]:
    for rel in ("manifest.json", "data/manifest.json"):
        p = extract_dir / rel
        if p.is_file():
            try:
                doc = json.loads(p.read_text(encoding="utf-8"))
            except Exception:
                return {}
            return doc if isinstance(doc, dict) else {}
    return {}


def _items_from_recovery_pending(extract_dir: Path) -> list[tuple[SourceIdentity, dict[str, Any], list[dict[str, Any]]]]:
    pending_path = extract_dir / "data" / "pending_ops.json"
    if not pending_path.is_file():
        return []
    try:
        doc = json.loads(pending_path.read_text(encoding="utf-8"))
    except Exception:
        return []
    ops = []
    if isinstance(doc, dict):
        ops = [o for o in _ensure_iterable(doc.get("operations"), name="operations") if isinstance(o, dict)]
    ctx_path = extract_dir / "data" / "context.json"
    ctx: dict[str, Any] = {}
    if ctx_path.is_file():
        try:
            raw = json.loads(ctx_path.read_text(encoding="utf-8"))
            if isinstance(raw, dict):
                ctx = raw
        except Exception:
            ctx = {}
    out = []
    for op in ops:
        op_type = str(op.get("op_type") or "")
        path = str(op.get("path") or "")
        if op_type != "save_results" and "save-results" not in path and "results" not in path:
            continue
        parsed = parse_pending_body(op.get("body"))
        payload = parsed.get("payload") if isinstance(parsed.get("payload"), dict) else {}
        rows = parsed.get("rows") if isinstance(parsed.get("rows"), list) else []
        src = SourceIdentity(
            exercise_id=_as_int(op.get("exercise_id") or ctx.get("exercise_id")),
            unit_id=_as_str(op.get("unit_id") or ctx.get("unit") or ctx.get("unit_key")),
            unit_name=_as_str(op.get("unit_name") or ctx.get("unit") or ctx.get("unit_name")),
            eval_item_id=_as_int(op.get("eval_item_id")),
            evaluation_list_id=_as_int(op.get("list_id")),
            evaluation_list_name=_as_str(op.get("list_name") or op.get("title")),
            evaluation_type=_as_str(op.get("evaluation_type") or "evaluation_list"),
            judge_id=_as_int(op.get("judge_id") or ctx.get("judge_id")),
            judge_name=_as_str(ctx.get("judge_name")),
            client_op_id=_as_str(parsed.get("client_op_id") or op.get("id")),
            device_id=_as_str(ctx.get("device_id")),
            stage_id=_as_str(op.get("stage_id") or payload.get("exercise_phase")),
            dilemma_name=_as_str(payload.get("dilemma_description")),
        )
        out.append((src, payload if payload else {"rows": rows}, rows))
    return out


def build_import_preview(
    db: Session,
    extract_dir: Path,
    *,
    package_hash: str,
    package_name: str,
    current_exercise_id: int | None,
) -> ImportPreview:
    applied = _package_applied(db, package_hash)
    previously = None
    if applied is not None:
        previously = {
            "applied_at": applied.applied_at.isoformat() if applied.applied_at else "",
            "restored_count": applied.restored_count,
        }
    warnings: list[str] = []
    items: list[ImportItemPreview] = []
    xlsx = _find_xlsx(extract_dir)
    recovered = _items_from_recovery_pending(extract_dir)
    manifest = _load_manifest(extract_dir)
    media_records = []
    if isinstance(manifest.get("media"), list):
        media_records = [m for m in manifest["media"] if isinstance(m, dict)]
    media_doc = extract_dir / "manifest.json"
    if not media_records and (extract_dir / "data" / "media_manifest.json").is_file():
        try:
            md = json.loads((extract_dir / "data" / "media_manifest.json").read_text(encoding="utf-8"))
            if isinstance(md, dict):
                media_records = [m for m in _ensure_iterable(md.get("records"), name="records") if isinstance(m, dict)]
        except Exception:
            pass

    sources: list[tuple[SourceIdentity, dict[str, Any], list[dict[str, Any]]]] = []
    if recovered:
        sources.extend(recovered)
    elif xlsx is not None:
        meta, rows, _list_name = _read_xlsx_bytes(xlsx.read_bytes())
        payload = {"rows": rows}
        sources.append((_source_from_meta(meta, rows), payload, rows))
        if meta.get("export_format_version") and meta.get("export_format_version") != EXPORT_FORMAT_VERSION:
            warnings.append(f"إصدار التصدير {meta.get('export_format_version')}")
    else:
        warnings.append("لا يوجد ملف Excel ولا عمليات حفظ في الحزمة.")

    pool = load_candidates(db, current_exercise_id=current_exercise_id)
    for src, payload, rows in sources:
        match = match_evaluation(
            db,
            src,
            current_exercise_id=current_exercise_id,
            candidates=pool,
        )
        media_n = 0
        eid = src.eval_item_id
        for rec in media_records:
            try:
                mid = int(rec.get("evaluation_list_item_id") or rec.get("eval_item_id") or 0)
            except (TypeError, ValueError):
                mid = 0
            if eid and mid == eid:
                media_n += 1
        dup = _client_op_seen(db, src.client_op_id)
        skip = ""
        if dup:
            skip = "نسخة مكررة عبر client_op_id"
        if previously and not skip:
            skip = "الحزمة مستوردة سابقاً"
            dup = True
        preview = ImportItemPreview(
            source=src,
            rows=rows,
            payload=payload if isinstance(payload, dict) else {"rows": rows},
            match=match,
            saved_at=_as_str((payload or {}).get("saved_at")),
            approved_at=_as_str((payload or {}).get("approved_at")),
            approved_source=bool((payload or {}).get("is_approved")),
            media_count=media_n,
            skip_reason=skip,
            duplicate=dup,
        )
        items.append(preview)

    return ImportPreview(
        package_hash=package_hash,
        package_name=package_name,
        extract_dir=str(extract_dir),
        previously_applied=previously,
        items=items,
        warnings=warnings,
        media_manifest=media_records,
    )


def apply_import(
    db: Session,
    preview: ImportPreview,
    *,
    selected_source_ids: list[int],
    manual_map: dict[int, int],
    operator_id: int | None,
    current_exercise_id: int | None,
    operation_type: str,
    force: bool = False,
    overwrite_approved: bool = False,
) -> dict[str, int]:
    from app.models import EvaluationListPdfItem, TabletClientOp
    from app.models.tablet_transfer import TabletImportPackage

    if preview.previously_applied and not force:
        raise RuntimeError("تم استيراد هذه الحزمة سابقاً. لن تُكرَّر النتائج.")

    wanted = {int(x) for x in selected_source_ids if int(x) > 0}
    if not wanted and manual_map:
        wanted = {int(k) for k in manual_map}
    restored = 0
    skipped = 0
    media_n = 0
    extract_dir = Path(preview.extract_dir)

    for item in preview.items:
        src_id = int(item.source.eval_item_id or 0)
        if wanted and src_id not in wanted and src_id:
            continue
        if item.duplicate and not force:
            skipped += 1
            continue
        match = item.match
        mapped_target = manual_map.get(src_id)
        if mapped_target:
            from app.models import EvaluationListPdfItem as EI

            target_item = db.get(EI, int(mapped_target))
            if target_item is None:
                skipped += 1
                continue
            saved = None
            from app.models import EvaluationListSavedResult

            saved = (
                db.query(EvaluationListSavedResult)
                .filter(EvaluationListSavedResult.evaluation_item_id == int(mapped_target))
                .order_by(EvaluationListSavedResult.id.desc())
                .first()
            )
            if saved is not None and (
                getattr(saved, "is_approved", False) or getattr(saved, "is_chief_approved", False)
            ):
                if not overwrite_approved:
                    skipped += 1
                    continue
            mapping_id = remember_mapping(
                db,
                item.source,
                CandidateTarget(
                    eval_item_id=int(mapped_target),
                    unit_id=getattr(target_item, "unit_level_key", "") or "",
                    unit_name=getattr(target_item, "unit_level_label", "") or "",
                    list_name=getattr(target_item, "text", "") or "",
                    evaluation_type="evaluation_list",
                    stage_id=getattr(target_item, "exercise_phase", "") or "",
                    stage_name=getattr(target_item, "exercise_phase", "") or "",
                    dilemma_id=None,
                    dilemma_name="",
                    exercise_id=getattr(target_item, "exercise_id", None),
                ),
                operator_id=operator_id,
                current_exercise_id=current_exercise_id,
            )
            target_eval_id = int(mapped_target)
            method = "MANUAL_MAPPING"
            status = "MATCHED_EXACT"
        else:
            if match.status == "CONFLICT" and not overwrite_approved:
                skipped += 1
                continue
            if match.status not in (
                "MATCHED_EXACT",
                "MATCHED_BY_UNIT_LIST",
                "MATCHED_BY_CONTEXT",
            ) or match.target is None:
                skipped += 1
                continue
            target_eval_id = int(match.target.eval_item_id)
            method = match.match_method
            status = match.status
            mapping_id = match.used_mapping_id
            target_item = db.get(EvaluationListPdfItem, target_eval_id)
            if target_item is None:
                skipped += 1
                continue

        payload = dict(item.payload or {})
        if "rows" not in payload:
            payload["rows"] = item.rows
        _write_saved_result(
            db,
            exercise_id=int(current_exercise_id or getattr(target_item, "exercise_id", 0) or 0),
            item=target_item,
            judge_id=item.source.judge_id,
            payload=payload,
        )
        if item.source.client_op_id and item.source.judge_id:
            existing = (
                db.query(TabletClientOp)
                .filter(
                    TabletClientOp.user_id == int(item.source.judge_id),
                    TabletClientOp.client_op_id == item.source.client_op_id,
                )
                .one_or_none()
            )
            if existing is None:
                db.add(
                    TabletClientOp(
                        user_id=int(item.source.judge_id),
                        exercise_id=current_exercise_id,
                        client_op_id=item.source.client_op_id[:120],
                        op_type="excel_import",
                        path="tablet-excel-import",
                        response_json="{}",
                    )
                )
        media_n += _import_media(
            db,
            extract_dir=extract_dir,
            exercise_id=int(current_exercise_id or 0),
            item=target_item,
            judge_id=item.source.judge_id,
            source_eval_id=src_id,
            target_eval_id=target_eval_id,
            records=preview.media_manifest,
        )
        record_transfer_audit(
            db,
            operation_type=operation_type,
            device_id=item.source.device_id,
            judge_id=item.source.judge_id,
            unit_id=item.source.unit_id,
            exercise_id=current_exercise_id,
            source_eval_item_id=src_id or None,
            target_eval_item_id=target_eval_id,
            match_method=method,
            match_status=status,
            mapping_id=mapping_id,
            operator_id=operator_id,
            package_hash=preview.package_hash,
            client_op_id=item.source.client_op_id,
        )
        restored += 1

    if restored <= 0 and skipped <= 0:
        raise RuntimeError("لا توجد عناصر قابلة للاستيراد.")
    db.add(
        TabletImportPackage(
            package_hash=preview.package_hash,
            package_name=preview.package_name[:256],
            operator_id=operator_id,
            restored_count=restored,
            details_json=json.dumps({"skipped": skipped, "media": media_n}, ensure_ascii=False),
        )
    )
    db.commit()
    return {"restored": restored, "skipped": skipped, "media": media_n}


def _import_media(
    db: Session,
    *,
    extract_dir: Path,
    exercise_id: int,
    item,
    judge_id: int | None,
    source_eval_id: int,
    target_eval_id: int,
    records: list[dict[str, Any]],
) -> int:
    from app.eval_criterion_media import (
        ALLOWED_PHOTO_TYPES,
        ALLOWED_VIDEO_TYPES,
        persist_criterion_medium,
    )
    from app.models import EvaluationCriterionMedia

    n = 0
    for rec in records:
        try:
            eid = int(rec.get("evaluation_list_item_id") or rec.get("eval_item_id") or 0)
        except (TypeError, ValueError):
            eid = 0
        if source_eval_id and eid and eid != source_eval_id:
            continue
        rel = str(rec.get("exported_path") or rec.get("path") or "").replace("\\", "/").lstrip("/")
        if not rel:
            continue
        src = extract_dir / rel
        if not src.is_file():
            alt = extract_dir / "media" / Path(rel).name
            src = alt if alt.is_file() else src
        if not src.is_file():
            continue
        client_uuid = str(rec.get("client_uuid") or rec.get("id") or "")
        if client_uuid:
            exists = (
                db.query(EvaluationCriterionMedia)
                .filter(EvaluationCriterionMedia.client_op_id == client_uuid[:120])
                .first()
            )
            if exists is not None:
                continue
        try:
            row_index = int(rec.get("row_index") or 0)
        except (TypeError, ValueError):
            row_index = 0
        kind = str(rec.get("media_kind") or "photo")
        mime = str(rec.get("mime_type") or "").split(";")[0].strip().lower()
        if "video" in kind.lower():
            if mime not in ALLOWED_VIDEO_TYPES:
                mime = "video/mp4"
        else:
            if mime not in ALLOWED_PHOTO_TYPES:
                mime = "image/jpeg"
        persist_criterion_medium(
            db,
            exercise_id=int(exercise_id),
            unit_level_key=getattr(item, "unit_level_key", "") or "",
            list_item_id=int(target_eval_id),
            bundle_action_eval_id=None,
            row_index=row_index,
            media_kind=kind,
            mime_type_in=mime,
            bin_data=src.read_bytes(),
            uploaded_by_id=judge_id,
            client_op_id=client_uuid,
        )
        n += 1
    return n

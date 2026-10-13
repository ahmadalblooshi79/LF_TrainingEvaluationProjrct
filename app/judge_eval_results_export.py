"""ملف Excel لنتائج قوائم التقييم كما يراها المحكم."""
from __future__ import annotations

import io

from app.evaluation_list_columns import display_grade_label, grade_label_from_percent
from app.evaluation_workflow import eval_judge_approved, eval_reopened_for_judge

KIND_PROCEDURE = "قوائم تقييم الإجراءات"
KIND_DILEMMA = "قوائم تقييم المعاضل"
POSITION_APPROVED = "معتمدة"
POSITION_PENDING = "بانتظار الاعتماد"
POSITION_EMPTY = "غير مدخلة"

_HEADERS = ("ت", "اسم القائمة", "النسبة المئوية", "النتيجة", "الموقف")


def list_result_fields(saved) -> dict[str, str]:
    """النسبة والنتيجة والموقف لصف واحد في ملف النتائج."""
    has_payload = bool(saved is not None and (getattr(saved, "payload_json", None) or "").strip())
    if not has_payload:
        return {"pct": "—", "grade": "—", "position": POSITION_EMPTY}
    pct_raw = getattr(saved, "total_pct", None)
    try:
        pct = float(pct_raw) if pct_raw is not None else None
    except (TypeError, ValueError):
        pct = None
    grade = display_grade_label(getattr(saved, "grade_label", "") or "")
    if not grade and pct is not None:
        grade = grade_label_from_percent(pct) or ""
    if eval_judge_approved(saved) and not eval_reopened_for_judge(saved):
        position = POSITION_APPROVED
    else:
        position = POSITION_PENDING
    pct_text = "—" if pct is None else f"{pct:.2f}%"
    return {"pct": pct_text, "grade": grade or "—", "position": position}


def build_judge_results_xlsx(*, kind_label: str, phases: list[dict]) -> bytes:
    """ورقة RTL: نوع القائمة، ثم المرحلة، ثم مسمى الوحدة واسم المحكم، ثم صفوف القوائم."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "النتائج"
    ws.sheet_view.rightToLeft = True
    thin = Border(
        left=Side(style="thin", color="C4B8A5"),
        right=Side(style="thin", color="C4B8A5"),
        top=Side(style="thin", color="C4B8A5"),
        bottom=Side(style="thin", color="C4B8A5"),
    )
    kind_fill = PatternFill("solid", fgColor="5C4033")
    phase_fill = PatternFill("solid", fgColor="8A6A45")
    unit_fill = PatternFill("solid", fgColor="F3E6D0")
    head_fill = PatternFill("solid", fgColor="5C4033")
    alt_fill = PatternFill("solid", fgColor="FBF6F0")
    pos_fills = {
        POSITION_APPROVED: PatternFill("solid", fgColor="E7F2E4"),
        POSITION_PENDING: PatternFill("solid", fgColor="F8E8C8"),
        POSITION_EMPTY: PatternFill("solid", fgColor="F4F0EA"),
    }
    kind_font = Font(name="Traditional Arabic", bold=True, color="FFFFFF", size=16)
    phase_font = Font(name="Traditional Arabic", bold=True, color="FFFFFF", size=13)
    unit_font = Font(name="Traditional Arabic", bold=True, color="5C4033", size=12)
    head_font = Font(name="Traditional Arabic", bold=True, color="FFFFFF", size=12)
    body_font = Font(name="Traditional Arabic", size=12)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    right = Alignment(horizontal="right", vertical="center", wrap_text=True)

    def paint(row: int, fill, font, align, values: tuple) -> None:
        for col, val in enumerate(values, start=1):
            cell = ws.cell(row, col, val)
            cell.fill = fill
            cell.font = font
            cell.alignment = align
            cell.border = thin
        ws.row_dimensions[row].height = 22

    def banner(row: int, text: str, fill, font) -> None:
        paint(row, fill, font, right, (text, "", "", "", ""))
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=5)
        ws.row_dimensions[row].height = 28

    row_ix = 1
    banner(row_ix, kind_label, kind_fill, kind_font)
    row_ix += 1
    wrote_any = False
    for phase in phases or []:
        blocks = [b for b in (phase.get("blocks") or []) if b.get("lists")]
        if not blocks:
            continue
        if wrote_any:
            row_ix += 1
        banner(
            row_ix,
            f"المرحلة: {(phase.get('phase_label') or 'بدون مرحلة').strip()}",
            phase_fill,
            phase_font,
        )
        row_ix += 1
        for block in blocks:
            unit_label = (block.get("unit_label") or "—").strip() or "—"
            judge_name = (block.get("judge_name") or "—").strip() or "—"
            paint(
                row_ix,
                unit_fill,
                unit_font,
                right,
                (f"مسمى الوحدة: {unit_label}", "", "", f"اسم المحكم: {judge_name}", ""),
            )
            ws.merge_cells(start_row=row_ix, start_column=1, end_row=row_ix, end_column=3)
            ws.merge_cells(start_row=row_ix, start_column=4, end_row=row_ix, end_column=5)
            row_ix += 1
            paint(row_ix, head_fill, head_font, center, _HEADERS)
            row_ix += 1
            for seq, item in enumerate(block.get("lists") or [], start=1):
                values = (
                    seq,
                    (item.get("title") or "").strip() or "—",
                    item.get("pct") or "—",
                    item.get("grade") or "—",
                    item.get("position") or POSITION_EMPTY,
                )
                for col, val in enumerate(values, start=1):
                    cell = ws.cell(row_ix, col, val)
                    cell.font = body_font
                    cell.alignment = right if col == 2 else center
                    cell.border = thin
                    if col == 5 and values[4] in pos_fills:
                        cell.fill = pos_fills[values[4]]
                    elif seq % 2 == 0:
                        cell.fill = alt_fill
                ws.row_dimensions[row_ix].height = 20
                row_ix += 1
            wrote_any = True
            row_ix += 1
    if not wrote_any:
        banner(row_ix, "لا توجد قوائم", unit_fill, unit_font)

    for idx, width in enumerate((8, 48, 18, 16, 22), start=1):
        ws.column_dimensions[get_column_letter(idx)].width = width
    ws.freeze_panes = "A2"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()

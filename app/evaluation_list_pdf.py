"""تصدير قائمة التقييم إلى PDF مطابق لنموذج Excel وبيانات النظام."""
from __future__ import annotations

import io
from pathlib import Path
from typing import Any

from app.xlsx_to_pdf import convert_xlsx_bytes_to_pdf

_FONT_REG = "EvalListAr"
_FONT_BOLD = "EvalListArBold"
_FONTS_REGISTERED = False


def build_evaluation_list_pdf_bytes(
    xlsx_bytes: bytes,
    *,
    signature_png: bytes | None = None,
) -> bytes:
    """
    يحوّل ملف Excel المعبأ إلى PDF (Excel ثم LibreOffice).
    إن تعذّر التحويل يُرسم النموذج بـ reportlab من نفس القيم.
    """
    pdf = convert_xlsx_bytes_to_pdf(xlsx_bytes)
    if pdf:
        return pdf
    return _reportlab_from_filled_xlsx(xlsx_bytes, signature_png=signature_png)


def _win_font(*names: str) -> Path | None:
    fonts = Path(r"C:\Windows\Fonts")
    for name in names:
        p = fonts / name
        if p.is_file():
            return p
    return None


def _ensure_fonts() -> None:
    global _FONTS_REGISTERED
    if _FONTS_REGISTERED:
        return
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    regular = _win_font("segoeui.ttf", "arial.ttf", "trado.ttf")
    bold = _win_font("segoeuib.ttf", "arialbd.ttf", "tradbdo.ttf") or regular
    if regular is None:
        raise RuntimeError("لم يُعثر على خط عربي مناسب في Windows Fonts.")
    pdfmetrics.registerFont(TTFont(_FONT_REG, str(regular)))
    pdfmetrics.registerFont(TTFont(_FONT_BOLD, str(bold)))
    _FONTS_REGISTERED = True


def _shape_ar(text: str) -> str:
    raw = (text or "").replace("\r\n", "\n")
    if not raw.strip():
        return ""
    try:
        import arabic_reshaper
        from bidi.algorithm import get_display

        return get_display(arabic_reshaper.reshape(raw))
    except Exception:
        return raw


def _display_cell(cell) -> str:
    v = getattr(cell, "value", None)
    if v is None:
        return ""
    nf = str(getattr(cell, "number_format", "") or "")
    if isinstance(v, (int, float)) and "%" in nf:
        pct = float(v) * 100.0
        if "0.00" in nf or "0.0%" in nf.replace(" ", ""):
            shown = round(pct * 100.0) / 100.0
            if shown == int(shown):
                return f"{int(shown)}%"
            return f"{shown}%"
        return f"{int(round(pct))}%"
    if isinstance(v, float) and v == int(v):
        return str(int(v))
    return str(v).replace("\xa0", " ").strip()


def _reportlab_from_filled_xlsx(
    xlsx_bytes: bytes,
    *,
    signature_png: bytes | None = None,
) -> bytes:
    from openpyxl import load_workbook
    from reportlab.lib.colors import HexColor, white
    from reportlab.lib.enums import TA_CENTER, TA_RIGHT
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import (
        Image as RLImage,
        Paragraph,
        SimpleDocTemplate,
        Spacer,
        Table,
        TableStyle,
    )

    _ensure_fonts()
    wb = load_workbook(io.BytesIO(xlsx_bytes), data_only=False)
    try:
        ws = wb.active
        mr = int(getattr(ws, "max_row", None) or 1)
        rows_out: list[list[Any]] = []
        title = _display_cell(ws.cell(1, 2))
        for r in range(2, mr + 1):
            element = _display_cell(ws.cell(r, 2))
            el_norm = " ".join((element or "").split())
            if "وصف المعضلة" in el_norm or "متطلبات تنفيذ" in el_norm:
                continue
            mx = _display_cell(ws.cell(r, 5))
            aq = _display_cell(ws.cell(r, 6))
            pct = _display_cell(ws.cell(r, 7))
            grade = _display_cell(ws.cell(r, 8))
            notes = _display_cell(ws.cell(r, 9))
            if not any((element, mx, aq, pct, grade, notes)):
                continue
            rows_out.append([notes, grade, pct, aq, mx, element])
    finally:
        wb.close()

    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf,
        pagesize=A4,
        rightMargin=12 * mm,
        leftMargin=12 * mm,
        topMargin=12 * mm,
        bottomMargin=12 * mm,
    )
    ink = HexColor("#1a1a2e")
    head_bg = HexColor("#c47a22")
    total_bg = HexColor("#f0d0a8")
    line = HexColor("#1f4e79")
    title_bg = HexColor("#f5e6c8")
    style_c = ParagraphStyle(
        "evalPdfC",
        fontName=_FONT_REG,
        fontSize=8,
        leading=11,
        alignment=TA_CENTER,
        textColor=ink,
    )
    style_r = ParagraphStyle(
        "evalPdfR",
        fontName=_FONT_REG,
        fontSize=8,
        leading=11,
        alignment=TA_RIGHT,
        textColor=ink,
    )
    style_title = ParagraphStyle(
        "evalPdfTitle",
        fontName=_FONT_BOLD,
        fontSize=11,
        leading=14,
        alignment=TA_CENTER,
        textColor=ink,
    )
    style_head = ParagraphStyle(
        "evalPdfHead",
        fontName=_FONT_BOLD,
        fontSize=8,
        leading=11,
        alignment=TA_CENTER,
        textColor=white,
    )

    def P(text: str, st=style_c):
        return Paragraph(_shape_ar(text or "").replace("\n", "<br/>"), st)

    story: list[Any] = []
    if title:
        story.append(
            Table(
                [[P(title, style_title)]],
                colWidths=[doc.width],
            )
        )
        story[-1].setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, -1), title_bg),
                    ("BOX", (0, 0), (-1, -1), 0.8, line),
                    ("TOPPADDING", (0, 0), (-1, -1), 8),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 8),
                ]
            )
        )
        story.append(Spacer(1, 4 * mm))

    headers = ["ملاحظات", "النتيجة", "النسبة", "المكتسبة", "القصوى", "عناصر التقييم"]
    data = [[P(h, style_head) for h in headers]]
    total_row_indexes: list[int] = []
    for i, row in enumerate(rows_out, start=1):
        blob = " ".join(str(x) for x in row)
        cells = [
            P(row[0], style_r),
            P(row[1], style_c),
            P(row[2], style_c),
            P(row[3], style_c),
            P(row[4], style_c),
            P(row[5], style_r),
        ]
        data.append(cells)
        if any(k in blob for k in ("إجمالي", "اجمالي", "النسبة العام", "التقدير العام")):
            total_row_indexes.append(i)

    col_w = [
        doc.width * 0.16,
        doc.width * 0.10,
        doc.width * 0.10,
        doc.width * 0.10,
        doc.width * 0.08,
        doc.width * 0.46,
    ]
    table = Table(data, colWidths=col_w, repeatRows=1)
    cmds = [
        ("BACKGROUND", (0, 0), (-1, 0), head_bg),
        ("TEXTCOLOR", (0, 0), (-1, 0), white),
        ("FONTNAME", (0, 0), (-1, 0), _FONT_BOLD),
        ("GRID", (0, 0), (-1, -1), 0.45, line),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("TOPPADDING", (0, 0), (-1, -1), 3),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
        ("LEFTPADDING", (0, 0), (-1, -1), 3),
        ("RIGHTPADDING", (0, 0), (-1, -1), 3),
    ]
    for ri in total_row_indexes:
        cmds.append(("BACKGROUND", (0, ri), (-1, ri), total_bg))
        cmds.append(("FONTNAME", (0, ri), (-1, ri), _FONT_BOLD))
    table.setStyle(TableStyle(cmds))
    story.append(table)

    if signature_png:
        try:
            img = RLImage(io.BytesIO(signature_png), width=42 * mm, height=16 * mm)
            img.hAlign = "LEFT"
            story.append(Spacer(1, 6 * mm))
            story.append(P("التوقيع", style_r))
            story.append(img)
        except Exception:
            pass

    doc.build(story)
    return buf.getvalue()
